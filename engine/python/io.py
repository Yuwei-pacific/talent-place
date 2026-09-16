#!/usr/bin/env python3
"""YAML and Excel helpers for poli-job-desk. JSON on stdin/stdout. PyYAML + openpyxl only."""
from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

COMPANY_HEADERS = [
    "Company / Outreach Account",
    "Brands / Business Units",
    "In Italy?",
    "Locations",
    "Strategic-fit Themes",
    "Matching Job Titles",
    "Job Links",
    "Role Count",
    "Curricular Evidence",
    "Work Modes",
    "Sources / Portals",
    "Previously Contacted?",
    "Contact Search Status",
    "Contact Name",
    "Contact Role",
    "Contact Email / LinkedIn",
    "Outreach Decision",
    "Notes",
    "Verification Status",
    "Last Checked",
]

ROLE_HEADERS = [
    "Company / Outreach Account",
    "Brand / Business Unit",
    "Job Title",
    "Location",
    "In Italy?",
    "Why It Fits",
    "Primary Job URL",
    "Alternate / Portal URLs",
    "Source / Portal",
    "Work Mode",
    "Curricular Status",
    "Previously Contacted?",
    "Verification Status",
    "Posted / Result Age",
    "Search Query",
    "Date Checked",
]

CONTACT_HEADERS = [
    "Company",
    "Email",
    "Phone",
    "First Contact Date",
    "Recall",
    "Notes",
]

SEARCH_LOG_HEADERS = [
    "Search Date",
    "Portal",
    "Search Query",
    "Filters",
    "Companies Added / Updated",
    "Jobs Selected",
    "Notes",
]


def emit(obj: object) -> None:
    json.dump(obj, sys.stdout, ensure_ascii=False)
    sys.stdout.write("\n")


def fail(message: str, code: int = 1) -> None:
    emit({"ok": False, "error": message})
    raise SystemExit(code)


def read_stdin_json() -> object:
    raw = sys.stdin.read()
    if not raw.strip():
        return {}
    return json.loads(raw)


def load_yaml(path: Path) -> object:
    import yaml

    text = path.read_text(encoding="utf-8")
    data = yaml.safe_load(text)
    return data


def dump_yaml(path: Path, data: object) -> None:
    import yaml

    path.parent.mkdir(parents=True, exist_ok=True)
    body = yaml.safe_dump(
        data,
        sort_keys=False,
        allow_unicode=True,
        default_flow_style=False,
        width=88,
    )
    path.write_text(body, encoding="utf-8")


def cmd_read_yaml(path: Path) -> None:
    if not path.exists():
        fail(f"missing {path}")
    emit({"ok": True, "path": str(path), "data": load_yaml(path)})


def cmd_write_yaml(path: Path) -> None:
    payload = read_stdin_json()
    if not isinstance(payload, dict) or "data" not in payload:
        fail("stdin must be JSON object with a data field")
    dump_yaml(path, payload["data"])
    emit({"ok": True, "path": str(path)})


def cmd_list_masters(root: Path) -> None:
    masters = []
    if root.exists():
        for child in sorted(root.iterdir()):
            profile = child / "profile.yaml"
            if child.is_dir() and profile.exists():
                try:
                    data = load_yaml(profile)
                    ident = data.get("id") if isinstance(data, dict) else None
                    official = data.get("official") if isinstance(data, dict) else None
                    name = official.get("name") if isinstance(official, dict) else None
                except Exception as orig_exc:  # noqa: BLE001
                    ident = None
                    name = None
                    error = str(orig_exc)
                else:
                    error = None
                pack = child / "views" / "search-pack.yaml"
                masters.append(
                    {
                        "id": child.name,
                        "profileId": ident,
                        "name": name or ident or child.name,
                        "hasCohort": (child / "cohort.yaml").exists(),
                        "hasPack": pack.exists(),
                        "error": error,
                    }
                )
    emit({"ok": True, "masters": masters})


def _headers_and_rows(ws) -> tuple[list, list[dict]]:
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        rows.append({str(headers[i] or ""): row[i] for i in range(len(headers))})
    return [h if h is not None else "" for h in headers], rows


def _header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for cell in ws[1]:
        if cell.value:
            mapping[str(cell.value)] = cell.column
    return mapping


def _ensure_sheet(wb, title: str, headers: list[str]):
    if title not in wb.sheetnames:
        ws = wb.create_sheet(title)
        for index, header in enumerate(headers, start=1):
            ws.cell(1, index, header)
        return ws
    ws = wb[title]
    have = _header_map(ws)
    next_col = max(have.values(), default=0) + 1
    for header in headers:
        if header not in have:
            ws.cell(1, next_col, header)
            have[header] = next_col
            next_col += 1
    return ws


def _write_cells(ws, row_index: int, values: dict) -> None:
    cols = _header_map(ws)
    for key, value in values.items():
        column = cols.get(key)
        if column is None:
            continue
        ws.cell(row_index, column, value)


def _next_row(ws) -> int:
    cols = _header_map(ws)
    key_col = next(iter(cols.values()), 1)
    last = 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, key_col).value not in (None, ""):
            last = row
    return last + 1


def _find_company_row(ws, name: str) -> int | None:
    cols = _header_map(ws)
    column = cols.get("Company / Outreach Account")
    if column is None:
        return None
    needle = name.strip().lower()
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, column).value
        if value and str(value).strip().lower() == needle:
            return row
    return None


def _cell(ws, row: int, header: str) -> str:
    cols = _header_map(ws)
    column = cols.get(header)
    if column is None:
        return ""
    value = ws.cell(row, column).value
    return "" if value is None else str(value)


def _parse_numbered(text: str) -> list[str]:
    import re

    items = []
    for line in (text or "").splitlines():
        line = re.sub(r"^\s*\d+\.\s*", "", line).strip()
        if line:
            items.append(line)
    if not items and (text or "").strip():
        items.append(text.strip())
    return items


def _format_numbered(items: list[str]) -> str:
    unique: list[str] = []
    seen: set[str] = set()
    for item in items:
        key = item.strip().rstrip("/").lower()
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(item.strip())
    return "\n".join(f"{index}. {item}" for index, item in enumerate(unique, start=1))


def _merge_semi(*parts: str) -> str:
    seen: list[str] = []
    keys: set[str] = set()
    for part in parts:
        for chunk in str(part or "").replace("\n", ";").split(";"):
            item = chunk.strip()
            key = item.lower()
            if item and key not in keys:
                keys.add(key)
                seen.append(item)
    return "; ".join(seen)


def _merge_lines(*parts: str) -> str:
    seen: list[str] = []
    keys: set[str] = set()
    for part in parts:
        for line in str(part or "").splitlines():
            item = line.strip()
            key = item.lower()
            if item and key not in keys:
                keys.add(key)
                seen.append(item)
    return "\n".join(seen)


def _curricular_summary(role_ws, company_name: str) -> str:
    cols = _header_map(role_ws)
    name_col = cols.get("Company / Outreach Account")
    status_col = cols.get("Curricular Status")
    if name_col is None:
        return "To verify"
    confirmed = 0
    to_verify = 0
    needle = company_name.strip().lower()
    for row in range(2, role_ws.max_row + 1):
        value = role_ws.cell(row, name_col).value
        if not value or str(value).strip().lower() != needle:
            continue
        status = "" if status_col is None else str(role_ws.cell(row, status_col).value or "").lower()
        if "confirm" in status:
            confirmed += 1
        else:
            to_verify += 1
    bits = []
    if confirmed:
        bits.append(f"Confirmed: {confirmed}")
    if to_verify:
        bits.append(f"To verify: {to_verify}")
    return "; ".join(bits) or "To verify"


def _italy_flag(*locations: str) -> str:
    import re

    text = " ".join(locations)
    has_it = bool(
        re.search(
            r"\b(italy|italia|milan|milano|rome|roma|turin|torino|bologna|florence|firenze|venice|venezia|naples|napoli|genoa|genova|parma|verona|padua|padova)\b",
            text,
            re.I,
        )
    )
    has_other = bool(
        re.search(
            r"\b(france|germany|spain|austria|netherlands|belgium|portugal|sweden|denmark|poland|ireland|greece|finland|munich|paris|berlin|madrid|barcelona|amsterdam|vienna|lisbon)\b",
            text,
            re.I,
        )
    )
    if has_it and has_other:
        return "Mixed"
    if has_it:
        return "Yes"
    if text.strip():
        return "No"
    return ""


def cmd_read_workbook(path: Path) -> None:
    if not path.exists():
        emit(
            {
                "ok": True,
                "path": str(path),
                "exists": False,
                "companies": [],
                "roles": [],
                "companyNames": [],
            }
        )
        return
    from openpyxl import load_workbook

    wb = load_workbook(path)
    companies = []
    roles = []
    if "Company Index" in wb.sheetnames:
        _, companies = _headers_and_rows(wb["Company Index"])
    if "Role Evidence" in wb.sheetnames:
        _, roles = _headers_and_rows(wb["Role Evidence"])
    names = []
    for row in companies:
        name = str(row.get("Company / Outreach Account") or "").strip()
        if name:
            names.append(name)
    emit(
        {
            "ok": True,
            "path": str(path),
            "exists": True,
            "companies": companies,
            "roles": roles,
            "companyNames": names,
        }
    )


def cmd_add_workbook(path: Path) -> None:
    payload = read_stdin_json()
    if not isinstance(payload, dict):
        fail("stdin must be a JSON object")
    companies = payload.get("companies") or []
    roles = payload.get("roles") or []
    if not isinstance(companies, list) or not isinstance(roles, list):
        fail("companies and roles must be arrays")

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        default = wb.active
        default.title = "Company Index"

    company_ws = _ensure_sheet(wb, "Company Index", COMPANY_HEADERS)
    role_ws = _ensure_sheet(wb, "Role Evidence", ROLE_HEADERS)
    _ensure_sheet(wb, "Previous Contact List", CONTACT_HEADERS)
    _ensure_sheet(wb, "Search Log", SEARCH_LOG_HEADERS)

    role_cols = _header_map(role_ws)
    url_col = role_cols.get("Primary Job URL")
    existing_urls = set()
    if url_col:
        for row in range(2, role_ws.max_row + 1):
            value = role_ws.cell(row, url_col).value
            if value:
                existing_urls.add(str(value).strip().rstrip("/").lower())

    added_companies = 0
    updated_companies = 0
    added_roles = 0
    skipped_roles = 0
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    touched: dict[str, list[dict]] = {}

    for role in roles:
        if not isinstance(role, dict):
            continue
        url = str(role.get("url") or role.get("Primary Job URL") or "").strip()
        title = str(role.get("title") or role.get("Job Title") or "").strip()
        company = str(role.get("company") or role.get("Company / Outreach Account") or "").strip()
        if not url or not title or not company:
            continue
        key = url.rstrip("/").lower()
        if key in existing_urls:
            skipped_roles += 1
            continue
        location = str(role.get("location") or role.get("Location") or "")
        source = str(role.get("source") or role.get("Source / Portal") or "")
        work_mode = str(role.get("workMode") or role.get("Work Mode") or "Physical location shown; onsite/hybrid status to verify")
        curricular = str(role.get("curricularStatus") or role.get("Curricular Status") or "To verify")
        verification = str(
            role.get("verificationStatus")
            or role.get("Verification Status")
            or "LinkedIn evidence captured; employer page, work mode and curricular status still require verification"
        )
        values = {
            "Company / Outreach Account": company,
            "Brand / Business Unit": str(role.get("brand") or role.get("Brand / Business Unit") or company),
            "Job Title": title,
            "Location": location,
            "In Italy?": str(role.get("inItaly") or role.get("In Italy?") or _italy_flag(location)),
            "Why It Fits": str(role.get("whyItFits") or role.get("Why It Fits") or ""),
            "Primary Job URL": url,
            "Alternate / Portal URLs": str(role.get("alternateUrls") or ""),
            "Source / Portal": source,
            "Work Mode": work_mode,
            "Curricular Status": curricular,
            "Previously Contacted?": str(role.get("previouslyContacted") or "No"),
            "Verification Status": verification,
            "Posted / Result Age": str(role.get("posted") or role.get("Posted / Result Age") or ""),
            "Search Query": str(role.get("query") or role.get("Search Query") or ""),
            "Date Checked": now,
        }
        row_index = _next_row(role_ws)
        _write_cells(role_ws, row_index, values)
        if url_col:
            cell = role_ws.cell(row_index, url_col)
            cell.hyperlink = url
            cell.font = Font(color="0563C1", underline="single")
        existing_urls.add(key)
        added_roles += 1
        touched.setdefault(company, []).append(values)

    company_meta = {
        str(item.get("company") or item.get("Company / Outreach Account") or "").strip(): item
        for item in companies
        if isinstance(item, dict)
    }

    for company_name, new_roles in touched.items():
        meta = company_meta.get(company_name) or {}
        row_index = _find_company_row(company_ws, company_name)
        creating = row_index is None
        if creating:
            row_index = _next_row(company_ws)
            added_companies += 1
        else:
            updated_companies += 1

        titles = _parse_numbered(_cell(company_ws, row_index, "Matching Job Titles"))
        links = _parse_numbered(_cell(company_ws, row_index, "Job Links"))
        for role in new_roles:
            title = str(role.get("Job Title") or "").strip()
            location = str(role.get("Location") or "").strip()
            label = f"{title} — {location}".strip(" —")
            if label:
                titles.append(label)
            url = str(role.get("Primary Job URL") or "").strip()
            if url:
                links.append(url)

        locations = _merge_lines(
            _cell(company_ws, row_index, "Locations"),
            str(meta.get("locations") or ""),
            *[str(role.get("Location") or "") for role in new_roles],
        )
        themes = _merge_semi(
            _cell(company_ws, row_index, "Strategic-fit Themes"),
            "; ".join(meta.get("themes") or []) if isinstance(meta.get("themes"), list) else str(meta.get("themes") or ""),
            *[str(role.get("Why It Fits") or "") for role in new_roles],
        )
        modes = _merge_lines(
            _cell(company_ws, row_index, "Work Modes"),
            str(meta.get("workMode") or ""),
            *[str(role.get("Work Mode") or "") for role in new_roles],
        )
        sources = _merge_semi(
            _cell(company_ws, row_index, "Sources / Portals"),
            *[str(role.get("Source / Portal") or "") for role in new_roles],
        )
        evidence = _curricular_summary(role_ws, company_name)

        values = {
            "Company / Outreach Account": company_name,
            "Brands / Business Units": _cell(company_ws, row_index, "Brands / Business Units")
            or str(meta.get("brand") or company_name),
            "In Italy?": _italy_flag(locations, _cell(company_ws, row_index, "In Italy?")),
            "Locations": locations,
            "Strategic-fit Themes": themes,
            "Matching Job Titles": _format_numbered(titles),
            "Job Links": _format_numbered(links),
            "Role Count": len(_parse_numbered(_format_numbered(links))),
            "Curricular Evidence": evidence,
            "Work Modes": modes,
            "Sources / Portals": sources,
            "Verification Status": _cell(company_ws, row_index, "Verification Status")
            or "LinkedIn evidence captured; employer page, work mode and curricular status still require verification",
            "Last Checked": now,
        }
        if creating:
            values.update(
                {
                    "Previously Contacted?": "No",
                    "Contact Search Status": "Not started",
                    "Contact Name": "",
                    "Contact Role": "",
                    "Contact Email / LinkedIn": "",
                    "Outreach Decision": "Review",
                    "Notes": "",
                }
            )
        _write_cells(company_ws, row_index, values)

    for ws, headers in ((company_ws, COMPANY_HEADERS), (role_ws, ROLE_HEADERS)):
        for index, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(index)].width = max(14, min(36, len(header) + 4))

    wb.save(path)
    emit(
        {
            "ok": True,
            "path": str(path),
            "addedCompanies": added_companies,
            "updatedCompanies": updated_companies,
            "addedRoles": added_roles,
            "skippedRoles": skipped_roles,
        }
    )


def main() -> int:
    if len(sys.argv) < 2:
        fail("usage: io.py <read-yaml|write-yaml|list-masters|read-workbook|add-workbook> [path]")
    cmd = sys.argv[1]
    path = Path(sys.argv[2]).expanduser() if len(sys.argv) > 2 else None
    if cmd == "read-yaml":
        if path is None:
            fail("read-yaml needs a path")
        cmd_read_yaml(path)
    elif cmd == "write-yaml":
        if path is None:
            fail("write-yaml needs a path")
        cmd_write_yaml(path)
    elif cmd == "list-masters":
        if path is None:
            fail("list-masters needs the workroot")
        cmd_list_masters(path)
    elif cmd == "read-workbook":
        if path is None:
            fail("read-workbook needs a path")
        cmd_read_workbook(path)
    elif cmd == "add-workbook":
        if path is None:
            fail("add-workbook needs a path")
        cmd_add_workbook(path)
    else:
        fail(f"unknown command {cmd}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as exc:  # noqa: BLE001
        fail(str(exc))
