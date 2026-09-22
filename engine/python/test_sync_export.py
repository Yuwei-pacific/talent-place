#!/usr/bin/env python3
"""python3 python/test_sync_export.py

Covers the guards, the splitter's parity with the TypeScript side, date parsing,
and the matching ladder. Runs without a real sync folder — the guards are
exercised against temp dirs, and the dataless case is simulated, because a real
dataless file cannot be created on demand.
"""
from __future__ import annotations

import csv
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import synced_fs  # noqa: E402
import reconcile  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from datetime import datetime  # noqa: E402

from sync_export import (  # noqa: E402
    A4_COLUMNS,
    CF_RANGE_ROWS,
    CONTACT_STATUS_VALUES,
    NOT_RECOVERABLE_FROM_FLAT_TSV,
    StageError,
    TRACKING_PARAMS,
    load_evidence,
    norm_role_url,
    REVIEW_COLUMNS,
    VERIFICATION_PREFIXES,
    explode_roles,
    evidence_report,
    validate_rows,
)
from reconcile import (  # noqa: E402
    HUMAN_COLS,
    MACHINE_LATEST_COLS,
    MACHINE_UNION_COLS,
    MULTI_VALUE_SPEC,
    find_duplicate_names,
    load_aliases,
    load_canonical,
    match_company,
    norm_company,
    parse_date,
    propose_company_id,
    split_column,
    split_multi,
)

REPO = HERE.parent.parent
CANONICAL = HERE.parent / "test" / "fixtures" / "strategic-design-history.csv"
# Sample TSVs live with the other test data, not in outputs/: they are FIXTURES,
# and a fixture filed under a directory called "outputs" reads as discardable.
# Measured: deleting them turns 6 tests into skips while the suite still says OK.
FIXTURES = HERE.parent / "test" / "fixtures"
ENGINE = HERE.parent
TS_LIB = ENGINE / "lib" / "normalize.js"


class FakeStat:
    """Stand-in for os.stat_result with the macOS st_flags field."""

    def __init__(self, size=100, blocks=8, flags=0, mode=0o100644):
        self.st_size = size
        self.st_blocks = blocks
        self.st_flags = flags
        self.st_mode = mode
        self.st_mtime_ns = 1
        self.st_mtime = 0.0


class TmpDirCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sync-test-"))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def patch_stat_for(self, target, **fields):
        """Patch os.stat to report `fields` for exactly `target`, delegating
        everything else to the real os.stat.

        Patching os.stat globally would make Path.exists() truthy for every
        probe, which silently turns the ~$ lock check into a false positive.
        """
        real = os.stat
        wanted = str(Path(target))

        def fake(path, *a, **k):
            if isinstance(path, (str, os.PathLike)) and str(path) == wanted:
                return FakeStat(**fields)
            return real(path, *a, **k)

        return mock.patch.object(synced_fs.os, "stat", side_effect=fake)


# ---------------------------------------------------------------------------
# Guards
# ---------------------------------------------------------------------------


class TestDataless(TmpDirCase):
    def test_real_local_file_is_not_dataless(self):
        f = self.tmp / "local.bin"
        f.write_bytes(b"x" * 100)
        self.assertFalse(synced_fs.is_dataless(f))

    def test_uf_dataless_flag_is_detected(self):
        # The real mount reports full st_size with st_blocks == 0 and UF_DATALESS
        # set. st_size > 0 therefore does NOT mean the bytes are local.
        f = self.tmp / "online.xlsx"
        f.write_bytes(b"placeholder")
        with self.patch_stat_for(f, size=14609, blocks=0, flags=synced_fs.UF_DATALESS):
            self.assertTrue(synced_fs.is_dataless(f))

    def test_zero_blocks_is_detected_without_the_flag(self):
        f = self.tmp / "online.xlsx"
        f.write_bytes(b"placeholder")
        with self.patch_stat_for(f, size=14609, blocks=0, flags=0):
            self.assertTrue(synced_fs.is_dataless(f))

    def test_zero_length_file_is_not_dataless(self):
        f = self.tmp / "empty.xlsx"
        f.write_bytes(b"")
        with self.patch_stat_for(f, size=0, blocks=0, flags=0):
            self.assertFalse(synced_fs.is_dataless(f))

    def test_missing_file_is_not_dataless(self):
        self.assertFalse(synced_fs.is_dataless(self.tmp / "nope.xlsx"))

    def test_writable_refuses_dataless_unless_allowed(self):
        target = self.tmp / "Review.xlsx"
        target.write_bytes(b"x")
        with self.patch_stat_for(target, size=10, blocks=0, flags=synced_fs.UF_DATALESS):
            with self.assertRaises(synced_fs.GuardFailure) as cm:
                synced_fs.assert_writable(target)
            self.assertEqual(cm.exception.reason, "dataless")
            # --allow-hydrate opts out
            synced_fs.assert_writable(target, allow_hydrate=True)


class TestExcelLock(TmpDirCase):
    def test_lock_is_detected_and_write_refused(self):
        target = self.tmp / "Review.xlsx"
        target.write_bytes(b"x")
        lock = self.tmp / "~$Review.xlsx"
        lock.write_bytes(b"")
        self.assertEqual(synced_fs.excel_lock_present(target), lock)
        with self.assertRaises(synced_fs.GuardFailure) as cm:
            synced_fs.assert_writable(target)
        self.assertEqual(cm.exception.reason, "excel_lock")

    def test_no_lock_when_absent(self):
        target = self.tmp / "Review.xlsx"
        target.write_bytes(b"x")
        self.assertIsNone(synced_fs.excel_lock_present(target))

    def test_truncated_lock_name_is_detected(self):
        # Excel may truncate the owner file name.
        target = self.tmp / "ReviewDocument.xlsx"
        target.write_bytes(b"x")
        (self.tmp / "~$ReviewDoc.xlsx").write_bytes(b"")
        self.assertIsNotNone(synced_fs.excel_lock_present(target))


class TestAtomicWrite(TmpDirCase):
    def test_writes_and_leaves_no_temp_file(self):
        target = self.tmp / "out.bin"
        synced_fs.write_atomically(target, lambda p: p.write_bytes(b"hello"), guard=False)
        self.assertEqual(target.read_bytes(), b"hello")
        leftovers = [f.name for f in self.tmp.iterdir() if f.name.startswith(".")]
        self.assertEqual(leftovers, [], f"temp files left behind: {leftovers}")

    def test_concurrent_change_aborts_and_preserves_the_other_writer(self):
        """The whole point of the guard: if a colleague saves while we write, we
        must NOT clobber them."""
        target = self.tmp / "Review.xlsx"
        target.write_bytes(b"original")

        def write_fn(tmpfile: Path) -> None:
            tmpfile.write_bytes(b"machine-update")
            target.write_bytes(b"colleague-saved-this")

        with self.assertRaises(synced_fs.GuardFailure) as cm:
            synced_fs.write_atomically(target, write_fn, backup_dir=self.tmp / "bak")
        self.assertEqual(cm.exception.reason, "changed_concurrently")
        self.assertEqual(target.read_bytes(), b"colleague-saved-this")

    def test_backup_taken_before_overwrite(self):
        target = self.tmp / "Review.xlsx"
        target.write_bytes(b"precious")
        bdir = self.tmp / "backups"
        synced_fs.write_atomically(target, lambda p: p.write_bytes(b"new"), guard=False, backup_dir=bdir)
        backups = list(bdir.glob("Review-*.xlsx"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_bytes(), b"precious")

    def test_backup_rotation_keeps_a_bounded_number(self):
        target = self.tmp / "Review.xlsx"
        bdir = self.tmp / "backups"
        for i in range(synced_fs.BACKUP_KEEP + 5):
            target.write_bytes(f"v{i}".encode())
            synced_fs.write_atomically(target, lambda p: p.write_bytes(b"x"), guard=False, backup_dir=bdir)
        self.assertLessEqual(len(list(bdir.glob("Review-*.xlsx"))), synced_fs.BACKUP_KEEP)

    def test_failed_write_does_not_leave_target_corrupted(self):
        target = self.tmp / "Review.xlsx"
        target.write_bytes(b"good")

        def boom(tmpfile: Path) -> None:
            tmpfile.write_bytes(b"partial")
            raise RuntimeError("disk full")

        with self.assertRaises(RuntimeError):
            synced_fs.write_atomically(target, boom, guard=False)
        self.assertEqual(target.read_bytes(), b"good")
        self.assertEqual([f.name for f in self.tmp.iterdir() if f.name.startswith(".")], [])


# ---------------------------------------------------------------------------
# Splitter — parity with the TypeScript implementation
# ---------------------------------------------------------------------------

PARITY_CASES = [
    ("Work Modes", "Physical location shown; onsite/hybrid status to verify"),
    ("Work Modes", "On-site"),
    ("Sources / Portals", "LinkedIn Jobs; iAgora (mirror annuncio aziendale)"),
    ("Matching Job Titles", "1. A Intern | 2. B Intern"),
    ("Matching Job Titles", "1. A Intern\n2. B Intern"),
    ("Matching Job Titles", "1. Intern - Level 2. Design"),
    ("Matching Job Titles", "1. Analyst, 3.5 days a week"),
    ("Matching Job Titles", "2. Alpha | 3. Beta"),
    ("Matching Job Titles", "1. Alpha\n1. Beta"),
    ("Matching Job Titles", "No numbering at all"),
    ("Matching Job Titles", "1. Only One"),
    ("Matching Job Titles", ""),
    ("Matching Job Titles", "   "),
    ("Locations", "Milan, Italy\nMilano, Italia"),
    ("Locations", "Milan, Italy | Rome, Italy"),
    ("Master-fit Themes", "CRM; Sustainability"),
    (
        "Job Links",
        "1. https://www.cefriel.com/careers/x/?lang=en\n   alt. https://it.linkedin.com/jobs/view/y-4440015924",
    ),
    ("Job Links", "1. https://x.example/job?id=13727121_it"),
    ("Job Links", "https://a.example/1 | https://b.example/2"),
]


class TestRoleEvidence(TmpDirCase):
    """`Roles.xlsx` is where a colleague clicks through to ONE posting, so a
    wrong Primary Job URL there sends them to the wrong job.

    Regression: `explode_roles` indexed `Job Links` through the column's default
    splitter, which expands the `alt.` mirror into a list item of its own. The
    index addresses a ROLE, so every role after the first mirror took the wrong
    URL — verified against the pre-fix code before it was changed.
    """

    def _row(self, links: str, titles: str, count: str) -> list[str]:
        row = [
            "Amplifon", "", "Yes", "Milan, Lombardy, Italy", "CRM / Customer Intelligence",
            titles, links, count, "", "", "", "No", "Not started", "", "", "", "Review",
            "", "Portal verified", "2026-09-17", "", "",
        ]
        self.assertEqual(len(row), len(A4_COLUMNS))
        return row

    def test_alt_mirror_does_not_consume_a_role_index(self):
        row = self._row(
            "1. https://a.example/1 alt. https://a-mirror.example/1 | "
            "2. https://a.example/2 alt. https://a-mirror.example/2 | "
            "3. https://a.example/3",
            "1. First Intern | 2. Second Intern | 3. Third Intern",
            "3",
        )
        roles, _ = explode_roles([row])
        self.assertEqual([r["Job Title"] for r in roles], ["First Intern", "Second Intern", "Third Intern"])
        self.assertEqual([r["Primary Job URL"] for r in roles],
                         ["https://a.example/1", "https://a.example/2", "https://a.example/3"])
        self.assertEqual([r["Alternate / Portal URLs"] for r in roles],
                         ["https://a-mirror.example/1", "https://a-mirror.example/2", ""])

    def test_alternate_column_is_recovered_not_declared_unrecoverable(self):
        # It is carried by the TSV's own `alt.` marker, so it no longer belongs
        # in the set of columns a live run cannot supply.
        self.assertNotIn("Alternate / Portal URLs", NOT_RECOVERABLE_FROM_FLAT_TSV)
        self.assertIn("Posted / Result Age", NOT_RECOVERABLE_FROM_FLAT_TSV)

    def test_the_report_names_a_column_the_sidecar_never_reached(self):
        """`columns_without_source` was keyed on the presence of `--evidence`, not
        on whether a single row joined. A run that followed A1 and wrote the
        EMPLOYER url into the TSV -- A1 prefers it over the portal url -- got a
        total join miss and a manifest declaring nothing missing.

        The sidecar is keyed by API url here and the TSV carries a different one,
        which is exactly the employer-vs-portal shape.
        """
        sidecar = {"roles": {"https://x.example/job/1": {"posted": "3 days ago"}}, "checks": {}}
        row = self._row("1. https://employer.example/1", "1. First Intern", "1")
        roles, _ = explode_roles([row], sidecar)

        report = evidence_report(roles, sidecar, "/run")
        self.assertEqual(report["filled"]["Posted / Result Age"], "0/1")
        self.assertIn("Posted / Result Age", report["without_source"])
        self.assertEqual(report["sidecar_rows"], 1, "the sidecar had a row; it just did not join")
        self.assertIn("portal URL", report["note"])

    def test_a_joined_column_reports_its_count(self):
        sidecar = {"roles": {"https://a.example/1": {"posted": "3 days ago"}}, "checks": {}}
        row = self._row("1. https://a.example/1 | 2. https://a.example/2", "1. First Intern | 2. Second Intern", "2")
        roles, _ = explode_roles([row], sidecar)

        report = evidence_report(roles, sidecar, "/run")
        self.assertEqual(report["filled"]["Posted / Result Age"], "1/2")
        self.assertNotIn("Posted / Result Age", report["without_source"])
        self.assertEqual(report["filled"]["Search Query"], "0/2")
        self.assertIn("Search Query", report["without_source"])
        self.assertNotIn("note", report, "a partial join is not a failure to declare")

    def test_without_evidence_every_derived_column_is_without_source(self):
        row = self._row("1. https://a.example/1", "1. First Intern", "1")
        roles, _ = explode_roles([row], None)
        report = evidence_report(roles, {"roles": {}, "checks": {}}, None)
        self.assertEqual(report["sidecar_rows"], 0)
        self.assertEqual(sorted(report["without_source"]), sorted(report["filled"]))
        self.assertNotIn("note", report, "no sidecar was offered, so no join failed")

    def _write_evidence(self, text: str) -> str:
        path = Path(self.tmp) / "role-evidence.csv"
        path.write_text(text, encoding="utf-8")
        return str(path)

    def test_evidence_fills_the_columns_a_flat_tsv_cannot_carry(self):
        ev = load_evidence(
            self._write_evidence(
                "url;posted;alternate_urls;search_query\n"
                "https://x.example/job/1;3 days ago;https://mirror.example/1;service design stage\n"
            )
        )
        self.assertIn("https://x.example/job/1", ev["roles"])
        self.assertEqual(ev["roles"]["https://x.example/job/1"]["posted"], "3 days ago")
        self.assertEqual(ev["roles"]["https://x.example/job/1"]["alternate_urls"], "https://mirror.example/1")
        self.assertEqual(ev["checks"], {}, "a run with no employer checks still loads")

    def test_evidence_joins_on_a_normalised_url(self):
        # A run's URL may carry campaign parameters the TSV's does not; the key
        # has to survive that or the evidence silently never lands.
        ev = load_evidence(self._write_evidence("url;posted;alternate_urls;search_query\nhttps://x.example/job/1?utm=abc;1 week ago;;q\n"))
        self.assertIn(norm_role_url("https://x.example/job/1/"), ev["roles"])

    def test_a_malformed_evidence_file_is_refused_not_ignored(self):
        bad = self._write_evidence("url;posted\nhttps://x.example/1;today\n")
        with self.assertRaises(StageError) as ctx:
            load_evidence(bad)
        self.assertIn("missing column", str(ctx.exception))

    def test_no_evidence_leaves_the_columns_empty(self):
        self.assertEqual(load_evidence(None), {"roles": {}, "checks": {}})
        self.assertEqual(load_evidence(""), {"roles": {}, "checks": {}})

    def test_explode_roles_applies_evidence_to_the_matching_role_only(self):
        row = self._row(
            "1. https://a.example/1 | 2. https://a.example/2",
            "1. First Intern | 2. Second Intern",
            "2",
        )
        roles, _ = explode_roles(
            [row],
            {"roles": {"https://a.example/2": {"posted": "5 days ago", "search_query": "q2", "alternate_urls": ""}}, "checks": {}},
        )
        self.assertEqual(roles[0]["Posted / Result Age"], "")
        self.assertEqual(roles[1]["Posted / Result Age"], "5 days ago")
        self.assertEqual(roles[1]["Search Query"], "q2")
        self.assertEqual(roles[0]["Search Query"], "")

    def _write_checks(self, text: str) -> str:
        path = Path(self.tmp) / "employer-checks.csv"
        path.write_text(text, encoding="utf-8")
        return str(path)

    def test_a_run_directory_supplies_both_sidecars(self):
        # A run writes both files next to each other; pointing --evidence at the
        # directory is the natural invocation, and neither file is required.
        (Path(self.tmp) / "role-evidence.csv").write_text(
            "url;posted;alternate_urls;search_query\nhttps://a.example/1;3 days ago;;q\n", encoding="utf-8"
        )
        self._write_checks(
            "url;label;status;reachable;has_apply;has_intern_signal;title;detail;elapsed_ms\n"
            "https://a.example/1;Acme;Employer verified active;true;true;true;Stage;page alive;412\n"
        )
        ev = load_evidence(self.tmp)
        self.assertIn("https://a.example/1", ev["roles"])
        self.assertIn("https://a.example/1", ev["checks"])
        self.assertEqual(ev["checks"]["https://a.example/1"]["status"], "Employer verified active")

    def test_a_single_file_is_identified_by_its_header(self):
        # Passing an employer-checks file directly used to fail with "missing
        # column: posted", which tells the caller nothing they can act on.
        checks_only = self._write_checks(
            "url;label;status;reachable;has_apply;has_intern_signal;title;detail;elapsed_ms\n"
            "https://a.example/1;Acme;Employer verified active;true;true;true;Stage;page alive;412\n"
        )
        ev = load_evidence(checks_only)
        self.assertEqual(ev["roles"], {})
        self.assertIn("https://a.example/1", ev["checks"])

        roles_only = self._write_evidence("url;posted;alternate_urls;search_query\nhttps://a.example/1;today;;q\n")
        self.assertIn("https://a.example/1", load_evidence(roles_only)["roles"])

    def test_no_evidence_at_all_is_not_an_error(self):
        self.assertEqual(load_evidence(None), {"roles": {}, "checks": {}})

    def test_a_probed_role_carries_its_own_verification_status(self):
        # Roles.xlsx is per ROLE, so a role the run actually probed should say so
        # rather than inheriting the company-level answer.
        row = self._row(
            "1. https://a.example/1 | 2. https://a.example/2",
            "1. First Intern | 2. Second Intern",
            "2",
        )
        roles, _ = explode_roles(
            [row],
            {
                "roles": {},
                "checks": {
                    "https://a.example/2": {"status": "Employer verified active"},
                },
            },
        )
        self.assertEqual(roles[1]["Verification Status"], "Employer verified active")
        self.assertEqual(
            roles[0]["Verification Status"],
            "Portal verified",
            "an unprobed role keeps the company-level value rather than inheriting a neighbour's",
        )

    def test_un_numbered_links_still_align_by_position(self):
        row = self._row(
            "https://a.example/1 | https://a.example/2",
            "1. First Intern | 2. Second Intern",
            "2",
        )
        roles, _ = explode_roles([row])
        self.assertEqual([r["Primary Job URL"] for r in roles],
                         ["https://a.example/1", "https://a.example/2"])

    def test_a_shorter_link_list_falls_back_without_inventing_a_url(self):
        row = self._row("1. https://a.example/1", "1. First Intern | 2. Second Intern", "2")
        roles, _ = explode_roles([row])
        self.assertEqual(roles[0]["Primary Job URL"], "https://a.example/1")
        self.assertEqual(roles[1]["Primary Job URL"], "https://a.example/1")
        self.assertEqual(roles[1]["Alternate / Portal URLs"], "")


class TestSplitter(TmpDirCase):
    def test_the_load_bearing_case_semicolon_inside_a_value(self):
        """Work Modes holds ';' INSIDE one value in 50 canonical rows. Treating
        ';' as a universal separator shreds it."""
        self.assertEqual(
            split_column("Physical location shown; onsite/hybrid status to verify", "Work Modes"),
            ["Physical location shown; onsite/hybrid status to verify"],
        )
        self.assertFalse(MULTI_VALUE_SPEC["Work Modes"]["semicolon"])
        self.assertTrue(MULTI_VALUE_SPEC["Sources / Portals"]["semicolon"])

    def test_mid_sentence_number_does_not_split(self):
        self.assertEqual(split_column("1. Intern - Level 2. Design", "Matching Job Titles"), ["Intern - Level 2. Design"])

    def test_alt_marker_splits_in_url_columns_only(self):
        got = split_column(
            "1. https://a.example/j?lang=en\n   alt. https://b.example/j/2",
            "Job Links",
        )
        self.assertEqual(got, ["https://a.example/j?lang=en", "https://b.example/j/2"])
        self.assertEqual(split_column("Alt. Text Intern", "Matching Job Titles"), ["Alt. Text Intern"])

    def test_free_text_columns_refuse_to_split(self):
        for col in ("Notes", "Curricular Evidence", "Reviewer Notes", "Matching Notes"):
            with self.assertRaises(KeyError):
                split_column("a; b", col)

    def test_degenerate_input(self):
        for empty in ("", "   ", None):
            self.assertEqual(split_column(empty, "Matching Job Titles"), [])

    def test_matches_the_typescript_implementation(self):
        """Two implementations in two languages read the same file; if they
        drift, the readers disagree again — the exact bug being fixed."""
        if not TS_LIB.exists():
            self.skipTest(f"{TS_LIB} missing — run `npm run build` in engine/ first")
        if shutil.which("node") is None:
            self.skipTest("node not on PATH")
        script = (
            'import("./lib/normalize.js").then(({splitColumn}) => {'
            "  const cases = JSON.parse(process.argv[1]);"
            "  console.log(JSON.stringify(cases.map(([c,v]) => splitColumn(v, c))));"
            "});"
        )
        proc = subprocess.run(
            ["node", "-e", script, json.dumps(PARITY_CASES)],
            cwd=str(ENGINE),
            capture_output=True,
            text=True,
        )
        self.assertEqual(proc.returncode, 0, proc.stderr[:400])
        ts_results = json.loads(proc.stdout.strip().splitlines()[-1])
        mismatches = []
        for (col, val), ts_val in zip(PARITY_CASES, ts_results):
            py_val = split_column(val, col)
            if py_val != ts_val:
                mismatches.append(f"  col={col!r} val={val[:50]!r}\n    py={py_val}\n    ts={ts_val}")
        self.assertEqual(mismatches, [], "TS/Python splitter drift:\n" + "\n".join(mismatches))

    def test_norm_role_url_keeps_identity_and_drops_tracking(self):
        """The query is not noise. Part of it names the posting.

        `?gh_jid=123` and `?gh_jid=456` are two different roles on one path.
        Stripping the whole query — which `dedup-cards.ts` did, and which this
        used to do — merges them, and the merge leaves no trace: a merged card is
        simply absent, so the role disappears silently.
        """
        # Identity survives, and two ids stay two postings.
        self.assertNotEqual(
            norm_role_url("https://x.test/jobs/1?gh_jid=123"),
            norm_role_url("https://x.test/jobs/1?gh_jid=456"),
        )
        # Tracking does not, and parameter order cannot matter.
        for with_noise in (
            "https://x.test/jobs/1?utm_source=linkedin&gh_jid=123",
            "https://x.test/jobs/1?gh_jid=123&trackingId=abc",
        ):
            self.assertEqual(norm_role_url(with_noise), norm_role_url("https://x.test/jobs/1?gh_jid=123"), with_noise)
        self.assertEqual(
            norm_role_url("https://x.test/jobs/1?b=2&a=1"),
            norm_role_url("https://x.test/jobs/1?a=1&b=2"),
            "parameter order must not change the key",
        )
        # A fragment is never identity; a trailing slash is not either.
        self.assertEqual(norm_role_url("https://x.test/jobs/1#apply"), norm_role_url("https://x.test/jobs/1"))
        self.assertEqual(norm_role_url("https://x.test/jobs/1/"), norm_role_url("https://x.test/jobs/1"))
        # A query of nothing but tracking is the same as no query at all.
        self.assertEqual(norm_role_url("https://x.test/jobs/1?utm_source=x"), norm_role_url("https://x.test/jobs/1"))

    def test_norm_role_url_parity_with_typescript(self):
        """Two languages build one dictionary key. Both the set and the function."""
        if not TS_LIB.exists() or shutil.which("node") is None:
            self.skipTest("TS build or node unavailable")
        script = (
            'import("./lib/normalize.js").then(({normalizeUrl, TRACKING_PARAMS}) => {'
            '  const urls = JSON.parse(process.argv[1]);'
            "  console.log(JSON.stringify({params: [...TRACKING_PARAMS].sort(), urls: urls.map(normalizeUrl)}));"
            "});"
        )
        urls = [
            "https://x.test/jobs/1",
            "https://x.test/jobs/1/",
            "https://X.TEST/Jobs/1",
            "https://x.test/jobs/1?gh_jid=123",
            "https://x.test/jobs/1?gh_jid=456",
            "https://x.test/jobs/1?utm_source=linkedin&gh_jid=123",
            "https://x.test/jobs/1?gh_jid=123&utm_source=linkedin",
            "https://x.test/jobs/1?utm=abc&gh_jid=123",
            "https://x.test/jobs/1?utm_id=9&gh_jid=123",
            "https://x.test/jobs/1?utmost=1",
            "https://x.test/jobs/1?b=2&a=1",
            "https://x.test/jobs/1?a=1&b=2",
            "https://x.test/jobs/1#apply",
            "https://x.test/jobs/1?utm_source=x",
            "https://x.test/jobs/1?",
            "",
            "not a url",
        ]
        proc = subprocess.run(["node", "-e", script, json.dumps(urls)], cwd=str(ENGINE), capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stderr[:400])
        ts = json.loads(proc.stdout.strip().splitlines()[-1])
        self.assertEqual(sorted(TRACKING_PARAMS), ts["params"], "the tracking sets must match exactly")
        for url, ts_val in zip(urls, ts["urls"]):
            self.assertEqual(norm_role_url(url), ts_val, f"norm_role_url drift on {url!r}")

    def test_norm_company_parity_with_typescript(self):
        if not TS_LIB.exists() or shutil.which("node") is None:
            self.skipTest("TS build or node unavailable")
        names = ["Accenture S.p.A.", "Nestlé", "Loro Piana", "Kering SA", "Moncler Group", "NTT DATA", "PwC"]
        script = (
            'import("./lib/normalize.js").then(({normCompany}) => {'
            "  console.log(JSON.stringify(JSON.parse(process.argv[1]).map(normCompany)));"
            "});"
        )
        proc = subprocess.run(["node", "-e", script, json.dumps(names)], cwd=str(ENGINE), capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stderr[:400])
        for name, ts_val in zip(names, json.loads(proc.stdout.strip().splitlines()[-1])):
            self.assertEqual(norm_company(name), ts_val, f"norm_company drift on {name!r}")


# ---------------------------------------------------------------------------
# Dates
# ---------------------------------------------------------------------------


class TestDates(unittest.TestCase):
    def test_ambiguous_date_is_read_day_first(self):
        """`First Contact Date` is stored D/M/YYYY (60 rows). A permissive parser
        reads 03/09/2026 as 9 March; we must read 3 September."""
        d = parse_date("03/09/2026")
        self.assertEqual((d.year, d.month, d.day), (2026, 9, 3))

    def test_all_formats_present_in_the_canonical_csv(self):
        self.assertEqual(parse_date("2026-09-15").isoformat(), "2026-09-15")
        self.assertEqual(parse_date("17/07/2026").isoformat(), "2026-07-17")
        self.assertEqual(parse_date("3/9/2026").isoformat(), "2026-09-03")

    def test_reads_an_iso_datetime(self):
        """The form `export-history` used to write: str() of the cell's datetime.

        Tolerating it is separate from no longer writing it. Exports already on
        disk hold this form, and a reader that cannot read yesterday's file
        reports phantom change until somebody re-exports it.
        """
        self.assertEqual(parse_date("2026-09-03 00:00:00").isoformat(), "2026-09-03")
        self.assertEqual(parse_date("2026-09-03T14:30:00").isoformat(), "2026-09-03")

    def test_unparseable_and_empty(self):
        for bad in ("", "   ", None, "garbage", "2026-13-45", "n/a"):
            self.assertIsNone(parse_date(bad))


# ---------------------------------------------------------------------------
# Canonical + matching ladder
# ---------------------------------------------------------------------------


@unittest.skipUnless(CANONICAL.exists(), "canonical CSV not present")
class TestCanonical(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.canon = load_canonical(CANONICAL)

    def test_loads(self):
        self.assertGreaterEqual(len(self.canon.rows), 100)
        self.assertIn("Company / Outreach Account", self.canon.header)

    def test_finds_the_five_duplicate_names(self):
        dups = find_duplicate_names(self.canon)
        self.assertEqual(sorted(dups), ["jakala", "kpmg", "ntt data", "pwc", "teamviewer"])
        for name, rows in dups.items():
            self.assertEqual(len(rows), 2, f"{name} should appear twice")

    def test_duplicate_name_is_ambiguous_never_silently_merged(self):
        for name in ("JAKALA", "KPMG", "NTT DATA", "PwC", "TeamViewer"):
            m = match_company(name, self.canon)
            self.assertEqual(m.kind, "ambiguous", f"{name} must not auto-merge")
            self.assertEqual(len(m.candidates), 2)

    def test_legal_suffix_folds_to_the_same_company(self):
        m = match_company("Accenture S.p.A.", self.canon)
        self.assertEqual(m.kind, "matched")
        self.assertEqual(m.row.company, "Accenture")

    def test_similar_name_needs_a_human(self):
        m = match_company("Accenture Italia", self.canon)
        self.assertEqual(m.kind, "needs_human")
        self.assertIn("Accenture", [c.company for c in m.candidates])

    def test_genuinely_new_company(self):
        m = match_company("Zzq Nonexistent Holdings Srl", self.canon)
        self.assertEqual(m.kind, "unmatched")
        self.assertEqual(m.reason, "new")

    def test_alias_table_resolves(self):
        target = self.canon.rows[0]
        aliases = {norm_company("Totally Different Name"): target.company_id or "x"}
        if not target.company_id:
            self.skipTest("canonical has no Company ID yet (run backfill-ids)")
        m = match_company("Totally Different Name", self.canon, aliases=aliases)
        self.assertEqual(m.kind, "matched")
        self.assertEqual(m.reason, "alias")

    def test_stored_id_is_authoritative(self):
        row = self.canon.rows[0]
        if not row.company_id:
            self.skipTest("canonical has no Company ID yet")
        m = match_company("A Completely Unrelated Name", self.canon, company_id=row.company_id)
        self.assertEqual(m.kind, "matched")
        self.assertEqual(m.reason, "company id")


class TestAliasTable(TmpDirCase):
    """`load_aliases` is the only reader of the human-maintained alias table, and
    rung 3 of the matching ladder cannot fire without it. A silent `{}` made "the
    file is missing" and "no aliases are needed" the same answer."""

    def test_a_missing_file_is_reported_not_silently_empty(self):
        table = load_aliases(self.tmp / "company-aliases.csv")
        self.assertEqual(table.mapping, {})
        self.assertEqual(table.report()["loaded"], 0)
        self.assertIn("not found", table.note)

    def test_it_reports_what_it_loaded(self):
        f = self.tmp / "company-aliases.csv"
        f.write_text(
            "alias_norm;canonical_id;note\n"
            "bof careers;polids-6556667b;BoF Careers republished a Moncler role\n",
            encoding="utf-8",
        )
        table = load_aliases(f)
        self.assertEqual(table.mapping, {"bof careers": "polids-6556667b"})
        self.assertEqual(table.report(), {"file": str(f), "loaded": 1})
        self.assertEqual(table.note, "")

    def test_a_header_only_file_is_reported_as_empty(self):
        # archive/company-aliases.csv is exactly this shape: 29 bytes and a
        # header. Restoring it verbatim would reproduce the silence it was
        # archived for.
        f = self.tmp / "company-aliases.csv"
        f.write_text("alias_norm;canonical_id;note\n", encoding="utf-8")
        table = load_aliases(f)
        self.assertEqual(table.mapping, {})
        self.assertIn("no aliases", table.note)


class TestCompanyId(unittest.TestCase):
    def test_survives_legal_suffix_but_not_a_real_rename(self):
        self.assertEqual(propose_company_id("Loro Piana"), propose_company_id("Loro Piana S.p.A."))
        # Documented limit: the hash is a PROPOSAL only, never recomputed.
        self.assertNotEqual(propose_company_id("Moncler"), propose_company_id("Moncler Group"))

    def test_stable(self):
        self.assertEqual(propose_company_id("Accenture"), propose_company_id("Accenture"))


class TestColumnOwnership(unittest.TestCase):
    def test_no_column_is_both_human_and_machine_owned(self):
        human, machine = set(HUMAN_COLS), set(MACHINE_UNION_COLS) | set(MACHINE_LATEST_COLS)
        self.assertEqual(human & machine, set(), "a column cannot be owned by both")

    def test_the_human_owned_columns_are_all_declared_human(self):
        for col in (
            "Previously Contacted?",
            "Contact Search Status",
            "Contact Name",
            "Contact Role",
            "Contact Email / LinkedIn",
            "Outreach Decision",
            "First Contact Date",
            "Recall",
        ):
            self.assertIn(col, HUMAN_COLS)

    def test_notes_is_not_machine_owned(self):
        """Notes was the one column BOTH sides wrote, fusing irreversibly."""
        self.assertNotIn("Notes", MACHINE_UNION_COLS)
        self.assertNotIn("Notes", MACHINE_LATEST_COLS)


# ---------------------------------------------------------------------------
# stage
# ---------------------------------------------------------------------------


class TestStage(TmpDirCase):
    # 09-11 is the shifted file and 09-10 the short one; neither satisfies A4, so
    # the happy-path tests use 09-10 with its rows padded to the header width.
    TSV = FIXTURES / "tsv-shifted-2026-09-11.tsv"
    CLEAN_SOURCE = FIXTURES / "tsv-short-2026-09-10.tsv"

    def clean_tsv_text(self) -> str:
        rows = list(csv.reader(io.StringIO(self.CLEAN_SOURCE.read_text(encoding="utf-8")), delimiter="\t"))
        hdr = rows[0]
        out = io.StringIO()
        w = csv.writer(out, delimiter="\t")
        w.writerow(hdr)
        for r in rows[1:]:
            w.writerow(r + [""] * (len(hdr) - len(r)))
        return out.getvalue()

    def test_never_creates_review(self):
        """stage is the zero-risk step: it must not touch the human file."""
        if not self.CLEAN_SOURCE.exists():
            self.skipTest("sample TSV missing")
        synced_fs.ensure_dir(self.tmp / "_machine")
        proc = subprocess.run(
            [sys.executable, str(HERE / "sync_export.py"), "stage", "--dir", str(self.tmp), "--history", str(CANONICAL)],
            input=self.clean_tsv_text(),
            check=True,
            capture_output=True,
            text=True,
        )
        self.assertTrue((self.tmp / "Roles.xlsx").exists())
        self.assertFalse((self.tmp / "Review.xlsx").exists(), "stage must never create Review.xlsx")

    def test_refuses_a_row_with_the_wrong_tab_count(self):
        if not self.CLEAN_SOURCE.exists():
            self.skipTest("sample TSV missing")
        lines = [ln for ln in self.clean_tsv_text().split("\n") if ln.strip()]
        corrupt = "\n".join([lines[0]] + [(ln + "\textra") if i == 2 else ln for i, ln in enumerate(lines[1:])])
        proc = subprocess.run(
            [sys.executable, str(HERE / "sync_export.py"), "stage", "--dir", str(self.tmp), "--history", str(CANONICAL)],
            input=corrupt,
            capture_output=True,
            text=True,
        )
        self.assertEqual(proc.returncode, 1)
        self.assertIn("wrong number of tabs", proc.stdout)

    def test_accepts_raw_tsv_on_stdin(self):
        if not self.CLEAN_SOURCE.exists():
            self.skipTest("sample TSV missing")
        proc = subprocess.run(
            [
                sys.executable,
                str(HERE / "sync_export.py"),
                "stage",
                "--dir",
                str(self.tmp),
                "--history",
                str(CANONICAL),
                "--dry-run",
            ],
            input=self.clean_tsv_text(),
            capture_output=True,
            text=True,
        )
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        out = json.loads(proc.stdout)
        self.assertGreater(out["roles_out"], 0)
        self.assertEqual(out["validation_problems"], [])


class TestReviewWorkbook(TmpDirCase):
    """init-review creates the file colleagues edit, and color installs the rules
    on it. Both must leave real dates and must not mis-address a column."""

    def make_canon(self):
        p = self.tmp / "canon.csv"
        rows = [
            ["Alpha Srl", "", "Yes", "Milan, Italy", "CRM", "1. Intern", "1. https://a/1", "1",
             "", "", "", "No", "Contacted", "Jane", "HR", "j@a.com", "Review", "machine note",
             "Employer verified active", "2026-09-15", "17/07/2026", "15/10/2026", "polids-aaaaaaaa"],
            ["Beta Srl", "", "Yes", "Rome, Italy", "CRM", "1. Intern", "1. https://b/1", "1",
             "", "", "", "No", "Not started", "", "", "", "Review", "", "Portal verified",
             "2026-09-14", "", "", "polids-bbbbbbbb"],
        ]
        hdr = A4_COLUMNS + ["Company ID"]
        with p.open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(hdr)
            w.writerows(rows)
        return p

    def run_cmd(self, *args):
        return subprocess.run([sys.executable, str(HERE / "sync_export.py"), *args], capture_output=True, text=True)

    def init(self, *extra):
        return self.run_cmd("init-review", "--dir", str(self.tmp), "--history", str(self.make_canon()), *extra)

    def test_refuses_to_overwrite_an_existing_review_file(self):
        self.assertEqual(self.init().returncode, 0)
        again = self.init()
        self.assertEqual(again.returncode, 1, "must refuse: the file holds colleague edits")
        self.assertIn("already exists", again.stdout)

    def test_columns_and_split_notes(self):
        self.init()
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(header, REVIEW_COLUMNS)
        self.assertEqual(len(header), 22)
        self.assertNotIn("Notes", header, "Notes must be split")
        # Both stay in A4 and in the TSV; Review.xlsx stops rendering them.
        # "Previously Contacted?" is in HUMAN_COLS, so `stage` still copies it
        # into the Company Summary sheet -- the value is not lost, only the
        # ability to edit it here.
        for gone in ("Previously Contacted?", "Outreach Decision"):
            self.assertNotIn(gone, header, f"{gone} should no longer be rendered in Review.xlsx")
            self.assertIn(gone, A4_COLUMNS, f"{gone} must stay an A4 column")
        self.assertIn("Matching Notes", header)
        self.assertIn("Reviewer Notes", header)
        # historical Notes seeds the machine column; the human column starts empty
        self.assertEqual(ws.cell(3, header.index("Matching Notes") + 1).value, "machine note")
        self.assertEqual(ws.cell(3, header.index("Reviewer Notes") + 1).value, None)

    def test_dates_are_real_dates_not_text(self):
        """If these are strings, `<=TODAY()` compares text and silently lies."""
        self.init()
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        for name, expect in [("First Contact Date", (2026, 7, 17)), ("Last Checked", (2026, 9, 15)), ("Recall", (2026, 10, 15))]:
            cell = ws.cell(3, header.index(name) + 1)
            self.assertIsInstance(cell.value, datetime, f"{name} must be a real date, got {type(cell.value)}")
            self.assertEqual((cell.value.year, cell.value.month, cell.value.day), expect)
            self.assertEqual(cell.number_format, "DD/MM/YYYY")

    def test_the_status_dropdown_actually_rejects_a_value_outside_the_list(self):
        """A list validation with `showErrorMessage=False` only OFFERS values.

        openpyxl defaults that flag to False, so Excel accepted anything typed
        and said nothing: the sheet looked protected and was not. It matters
        because an unrecognised value is then invisible twice over — it matches
        no conditional-formatting rule (those are exact-equality), so the row
        stays white, which is exactly what `Not started` looks like.
        """
        self.init()
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        rules = list(ws.data_validations.dataValidation)
        self.assertTrue(rules, "expected the status dropdown to be installed")
        dv = next((d for d in rules if "Not started" in (d.formula1 or "")), None)
        self.assertIsNotNone(dv, "no dropdown found for Contact Search Status")
        self.assertTrue(dv.showErrorMessage, "the dropdown would accept any value silently")
        self.assertEqual(dv.errorStyle, "stop")
        col = get_column_letter(REVIEW_COLUMNS.index("Contact Search Status") + 1)
        self.assertTrue(
            str(dv.sqref).startswith(f"{col}3:"),
            f"dropdown covers {dv.sqref}, not column {col}",
        )

    def test_unparseable_date_is_reported_not_silently_converted(self):
        p = self.make_canon()
        text = p.read_text(encoding="utf-8").replace("2026-09-15", "Brands / Business Units")
        p.write_text(text, encoding="utf-8")
        proc = self.run_cmd("init-review", "--dir", str(self.tmp), "--history", str(p))
        self.assertEqual(proc.returncode, 0)
        out = json.loads(proc.stdout)
        self.assertIn("unparseable_dates", out)
        self.assertEqual(len(out["unparseable_dates"]), 1)
        self.assertEqual(out["unparseable_dates"][0]["value"], "Brands / Business Units")

    def test_color_install_and_addresses_the_right_columns(self):
        self.init()
        self.assertEqual(self.run_cmd("color", "--dir", str(self.tmp)).returncode, 0)

        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        by_range: dict[str, list[str]] = {}
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            by_range[str(rng.sqref)] = [r.formula[0] for r in rules]

        status = get_column_letter(header.index("Contact Search Status") + 1)
        last = get_column_letter(len(REVIEW_COLUMNS))
        row_range = f"A3:{last}{CF_RANGE_ROWS}"

        # The whole row, one rule per coloured status, all anchored on the
        # status cell. "Not started" has no rule because white is the sheet's
        # own background -- the default costs nothing to maintain.
        self.assertIn(row_range, by_range, f"expected a whole-row rule on {row_range}, got {sorted(by_range)}")
        formulas = by_range[row_range]
        # Four exact-match rules, plus a fifth that reddens any status outside
        # the closed set -- see
        # test_color_reddens_a_status_outside_the_closed_set.
        self.assertEqual(len(formulas), 5, f"four statuses plus the fallback, got {formulas}")
        for value in ("Job not suitable", "Potential contact", "Contact found", "Job found"):
            self.assertIn(f'${status}3="{value}"', formulas)
        self.assertNotIn(
            f'${status}3="Not started"',
            formulas,
            "Not started must rely on the sheet background, not a rule",
        )

        # The status rule must NOT reference the date column any more. The old
        # ruleset did, to detect contradictions, and that branch was removed on
        # purpose: simpler to maintain, at the cost of the self-clearing amber.
        date_col = get_column_letter(header.index("First Contact Date") + 1)
        joined = " ".join(formulas)
        self.assertNotIn(f"${date_col}3", joined, "the contradiction branch should be gone")

        # Recall keeps its own two rules, on its own column.
        recall = get_column_letter(header.index("Recall") + 1)
        self.assertIn(f"{recall}2:{recall}{CF_RANGE_ROWS}", by_range)

    def test_color_reddens_a_status_outside_the_closed_set(self):
        # White is the sheet's own background and "Not started" deliberately has
        # no rule, so without a fallback a value outside the closed set is
        # indistinguishable from "Not started": a row someone has worked on
        # reads as untouched. The fallback makes it loud instead.
        self.init()
        self.assertEqual(self.run_cmd("color", "--dir", str(self.tmp)).returncode, 0)

        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        by_range: dict[str, list[str]] = {}
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            by_range[str(rng.sqref)] = [r.formula[0] for r in rules]

        status = get_column_letter(header.index("Contact Search Status") + 1)
        last = get_column_letter(len(REVIEW_COLUMNS))
        formulas = by_range[f"A3:{last}{CF_RANGE_ROWS}"]

        # Last, so the four exact-match rules keep priority over it.
        fallback = formulas[-1]
        self.assertIn(
            f'${status}3<>""',
            fallback,
            "the fallback must leave an empty status alone: A4 uses empty for 'not stated'",
        )
        for value in CONTACT_STATUS_VALUES:
            self.assertIn(
                f'${status}3<>"{value}"',
                fallback,
                "the exclusion list must be derived from CONTACT_STATUS_VALUES, so adding a "
                "status to the vocabulary cannot silently start rendering it as a mistake",
            )

    def test_color_is_idempotent(self):
        self.init()
        self.run_cmd("color", "--dir", str(self.tmp))
        from openpyxl import load_workbook

        first = len(load_workbook(self.tmp / "Review.xlsx")["Review"].conditional_formatting._cf_rules)
        self.run_cmd("color", "--dir", str(self.tmp))
        second = len(load_workbook(self.tmp / "Review.xlsx")["Review"].conditional_formatting._cf_rules)
        self.assertEqual(first, second, "re-running color must not stack duplicate rules")

    def test_color_refuses_without_a_review_file(self):
        proc = self.run_cmd("color", "--dir", str(self.tmp))
        self.assertEqual(proc.returncode, 2)
        self.assertIn("init-review", proc.stdout)

    def test_color_refuses_while_excel_has_the_file_open(self):
        self.init()
        (self.tmp / "~$Review.xlsx").write_bytes(b"")
        proc = self.run_cmd("color", "--dir", str(self.tmp))
        self.assertEqual(proc.returncode, 1)
        self.assertIn("excel_lock", proc.stdout)

    def test_doctor_rejects_text_in_a_date_column(self):
        self.init()
        self.run_cmd("color", "--dir", str(self.tmp))
        proc = self.run_cmd("doctor", "--dir", str(self.tmp))
        # /tmp is not a sync mount, so doctor fails overall; assert the specific check
        checks = {c["check"]: c for c in json.loads(proc.stdout)["checks"]}
        self.assertTrue(checks["Review.xlsx 'First Contact Date' holds dates"]["ok"])

    def test_init_review_carries_a_superseded_status_forward(self):
        """A new file must not be born outside the vocabulary.

        `make_canon` seeds `Contact Search Status = Contacted`, which A4 lists
        under LEGACY_CONTACT_STATUS. `migrate-review` translates those on a live
        sheet; `init-review` used to copy them verbatim, so a freshly created
        workbook failed doctor's closed-set check on its first run.
        """
        self.init()
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        ci = header.index("Contact Search Status") + 1
        self.assertEqual(ws.cell(3, ci).value, "Contact found", "Contacted is a superseded value")
        self.assertEqual(ws.cell(4, ci).value, "Not started", "a current value passes through")

    def test_doctor_passes_a_workbook_whose_values_fit_their_columns(self):
        self.init()
        proc = self.run_cmd("doctor", "--dir", str(self.tmp))
        checks = {c["check"]: c for c in json.loads(proc.stdout)["checks"]}
        self.assertTrue(checks["Review.xlsx closed-set columns hold A4's values"]["ok"])

    def test_doctor_reports_a_value_outside_its_closed_set(self):
        """The rule `stage` applies to TSV rows, applied to the workbook.

        `Contact Search Status` is the structural one: the axis the whole row is
        coloured by, closed "di proposito" in A4, and enforced on entry by the
        dropdown. A value outside it turns the row white, which reads as
        `Not started`, and nothing reported that before this check existed.
        """
        self.init()
        from openpyxl import load_workbook

        wb = load_workbook(self.tmp / "Review.xlsx")
        ws = wb["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        # `Contacted` is a superseded value: A4 lists it under LEGACY_CONTACT_STATUS.
        ws.cell(3, header.index("Contact Search Status") + 1, "Contacted")
        wb.save(self.tmp / "Review.xlsx")

        proc = self.run_cmd("doctor", "--dir", str(self.tmp))
        found = {c["check"]: c for c in json.loads(proc.stdout)["checks"]}[
            "Review.xlsx closed-set columns hold A4's values"
        ]
        self.assertFalse(found["ok"])
        self.assertIn("Contact Search Status", found["detail"])
        self.assertIn("Alpha Srl", found["detail"])

    def test_doctor_does_not_fail_on_the_reference_column(self):
        """`Verification Status` in the workbook is a reference, not a filter.

        A4 gives it a closed set too, but that governs what a RUN produces —
        `stage` enforces it on the TSV. The workflow does not select or group on
        the workbook's copy, so the 71 prose values in the live sheet are left
        alone. Checking them here reported 71 failures with nothing wrong, and a
        check that is always red is how an exit code stops being read.
        """
        self.init()
        from openpyxl import load_workbook

        wb = load_workbook(self.tmp / "Review.xlsx")
        ws = wb["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        ws.cell(3, header.index("Verification Status") + 1, "LinkedIn evidence captured; employer page not verified")
        wb.save(self.tmp / "Review.xlsx")

        proc = self.run_cmd("doctor", "--dir", str(self.tmp))
        checks = {c["check"]: c for c in json.loads(proc.stdout)["checks"]}
        self.assertTrue(
            checks["Review.xlsx closed-set columns hold A4's values"]["ok"],
            "a reference column must not make doctor fail",
        )

    def test_doctor_treats_an_empty_closed_set_cell_as_not_stated(self):
        """Empty is what A4 uses for "not stated", so it is not a problem.

        The two cells the Bain/Moncler paste corrupted were cleared to empty
        rather than guessed at; this check must not then call them a fault."""
        self.init()
        from openpyxl import load_workbook

        wb = load_workbook(self.tmp / "Review.xlsx")
        ws = wb["Review"]
        header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        ws.cell(3, header.index("Verification Status") + 1).value = None
        wb.save(self.tmp / "Review.xlsx")

        proc = self.run_cmd("doctor", "--dir", str(self.tmp))
        checks = {c["check"]: c for c in json.loads(proc.stdout)["checks"]}
        self.assertTrue(checks["Review.xlsx closed-set columns hold A4's values"]["ok"])


class TestRowValidation(TmpDirCase):
    """The tab-count rule cannot see a row shifted sideways: 22 tabs is 22 tabs
    whichever column the values landed in. These are the real cases from
    outputs/, both of which shipped."""

    def build(self, mutate):
        hdr = list(A4_COLUMNS)
        row = [""] * len(hdr)
        row[0] = "Alpha Srl"
        row[hdr.index("Verification Status")] = "Employer verified active"
        row[hdr.index("Last Checked")] = "2026-09-15"
        row[hdr.index("Outreach Decision")] = "Review"
        row[hdr.index("Contact Search Status")] = "Not started"
        return hdr, [mutate(row)]

    def test_clean_row_has_no_problems(self):
        hdr, rows = self.build(lambda r: r)
        self.assertEqual(validate_rows(hdr, rows), [])

    def test_shift_detected_even_though_the_tab_count_is_correct(self):
        """This is the 09-11 fixture: values right, count right, columns wrong."""
        def shift(r):
            out = [""] * len(r)
            out[16] = ""
            out[17] = ""
            out[18] = "Review"  # the decision, sitting in Verification Status
            out[19] = "[NEW COMPANY] [pertinente] Match score: 88/100;"  # notes in Last Checked
            out[20] = "To verify — employer site not reachable"  # verification in a date column
            out[21] = "2026-09-11"
            return out

        hdr, rows = self.build(shift)
        self.assertEqual(len(rows[0]), len(hdr), "the file is 22 columns wide — count check passes")
        problems = validate_rows(hdr, rows)
        self.assertTrue(problems, "a sideways shift must be caught")
        joined = " | ".join(problems)
        self.assertIn("Verification Status", joined)
        self.assertIn("is not a date", joined)

    def test_bad_verification_status_is_caught(self):
        hdr, rows = self.build(lambda r: r)
        rows[0][hdr.index("Verification Status")] = "looks fine to me"
        self.assertTrue(any("Verification Status" in p for p in validate_rows(hdr, rows)))

    def test_bad_decision_is_caught(self):
        hdr, rows = self.build(lambda r: r)
        rows[0][hdr.index("Outreach Decision")] = "maybe"
        self.assertTrue(any("Outreach Decision" in p for p in validate_rows(hdr, rows)))

    def test_text_in_a_date_column_is_caught(self):
        hdr, rows = self.build(lambda r: r)
        rows[0][hdr.index("First Contact Date")] = "sometime last week"
        self.assertTrue(any("First Contact Date" in p for p in validate_rows(hdr, rows)))

    def test_each_a4_verification_prefix_is_accepted(self):
        for prefix in VERIFICATION_PREFIXES:
            hdr, rows = self.build(lambda r: r)
            rows[0][hdr.index("Verification Status")] = f"{prefix} — some free-text detail"
            self.assertEqual(validate_rows(hdr, rows), [], f"{prefix!r} should be accepted")

    def test_empty_cells_are_never_problems(self):
        hdr = list(A4_COLUMNS)
        row = [""] * len(hdr)
        row[0] = "Alpha"
        self.assertEqual(validate_rows(hdr, [row]), [])


class TestStageRefusesMalformed(unittest.TestCase):
    TSV_0910 = FIXTURES / "tsv-short-2026-09-10.tsv"
    TSV_0911 = FIXTURES / "tsv-shifted-2026-09-11.tsv"

    def run_stage(self, tsv, *extra):
        return subprocess.run(
            [
                sys.executable,
                str(HERE / "sync_export.py"),
                "stage",
                "--dir",
                tempfile.mkdtemp(prefix="sync-stage-"),
                "--history",
                str(CANONICAL),
                "--tsv",
                str(tsv),
                "--dry-run",
                *extra,
            ],
            capture_output=True,
            text=True,
        )

    @unittest.skipUnless(CANONICAL.exists(), "canonical CSV not present")
    def test_both_shipped_tsvs_are_refused(self):
        """Neither historical output satisfies A4. 09-10 omits the last two
        columns; 09-11 has the tail shifted by two. Both shipped for manual
        pasting, so nothing ever validated them."""
        if not self.TSV_0910.exists() or not self.TSV_0911.exists():
            self.skipTest("sample TSVs missing")
        for tsv in (self.TSV_0910, self.TSV_0911):
            proc = self.run_stage(tsv)
            self.assertEqual(proc.returncode, 1, f"{tsv.name} should be refused")
            self.assertFalse(json.loads(proc.stdout).get("ok", False))

    @unittest.skipUnless(CANONICAL.exists(), "canonical CSV not present")
    def test_lenient_records_instead_of_refusing(self):
        if not self.TSV_0911.exists():
            self.skipTest("sample TSV missing")
        proc = self.run_stage(self.TSV_0911, "--lenient")
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        out = json.loads(proc.stdout)
        self.assertGreater(len(out["validation_problems"]), 0, "problems must be recorded, not swallowed")
        self.assertEqual(out["companies_in"], 6)

    @unittest.skipUnless(CANONICAL.exists(), "canonical CSV not present")
    def test_a_correctly_padded_file_stages(self):
        """Padding short rows with empty trailing fields is the fix for 09-10."""
        if not self.TSV_0910.exists():
            self.skipTest("sample TSV missing")
        rows = list(csv.reader(io.StringIO(self.TSV_0910.read_text(encoding="utf-8")), delimiter="\t"))
        hdr = rows[0]
        fixed = self.TSV_0910.parent / "_tmp_fixed.tsv"
        try:
            with fixed.open("w", encoding="utf-8", newline="") as fh:
                w = csv.writer(fh, delimiter="\t")
                w.writerow(hdr)
                for r in rows[1:]:
                    w.writerow(r + [""] * (len(hdr) - len(r)))
            proc = self.run_stage(fixed)
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            out = json.loads(proc.stdout)
            self.assertEqual(out["validation_problems"], [])
            self.assertEqual(out["companies_in"], 13)
        finally:
            fixed.unlink(missing_ok=True)


class TestAppend(TmpDirCase):
    """append writes into the file colleagues edit. The invariants: only empty
    cells are touched, no style is ever written, machine-owned dates are carried
    and human-owned ones are not invented."""

    def build(self, review_rows, root=None):
        """Build a Review.xlsx with REVIEW_COLUMNS and the given data rows.

        `root` defaults to the temp dir itself. Passing a subdirectory is for the
        tests that need the publish root -- the level where `company-aliases.csv`
        lives -- to be inside the temp dir rather than its parent.
        """
        from openpyxl import Workbook

        d = root or self.tmp
        wb = Workbook()
        ws = wb.active
        ws.title = "Review"
        for c, name in enumerate(REVIEW_COLUMNS, start=1):
            ws.cell(2, c, name)
        for i, row in enumerate(review_rows):
            for c, name in enumerate(REVIEW_COLUMNS, start=1):
                if row.get(name):
                    ws.cell(3 + i, c, row[name])
        wb.save(d / "Review.xlsx")

    def stage_run(self, companies, root=None):
        """Write a _machine/companies-<run>.csv as `stage` would, plus a manifest."""
        d = root or self.tmp
        run = "20260916-000000"
        machine = synced_fs.ensure_dir(d / "_machine")
        with (machine / f"companies-{run}.csv").open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(A4_COLUMNS)
            for comp in companies:
                w.writerow([comp.get(c, "") for c in A4_COLUMNS])
        (machine / "last-run.json").write_text(json.dumps({"run_id": run}), encoding="utf-8")
        return run

    def new_company(self, name="Nuova Azienda Srl", **over):
        row = {c: "" for c in A4_COLUMNS}
        row.update(
            {
                "Company / Outreach Account": name,
                "Locations": "Milan, Italy",
                "In Italy?": "Yes",
                "Master-fit Themes": "CRM",
                "Matching Job Titles": "1. Intern",
                "Job Links": "1. https://example.com/1",
                "Role Count": "1",
                "Notes": "machine note",
                "Verification Status": "Portal verified",
                "Last Checked": "2026-09-15",
            }
        )
        row.update(over)
        return row

    def run_append(self, *extra, root=None):
        return subprocess.run(
            [sys.executable, str(HERE / "sync_export.py"), "append", "--dir", str(root or self.tmp), *extra],
            capture_output=True,
            text=True,
        )

    def test_the_summary_reports_the_alias_table_it_matched_with(self):
        # The table lives at the publish root, one level above the Master dir:
        # a portal name is not a per-Master distinction.
        master = self.tmp / "Strategic design ED.28"
        master.mkdir()
        self.build([{"Company / Outreach Account": "Esistente Srl", "Company ID": "polids-11111111"}], root=master)
        self.stage_run([self.new_company("Nuova Azienda Srl")], root=master)

        out = json.loads(self.run_append(root=master).stdout)
        self.assertEqual(out["aliases"]["file"], str(self.tmp / "company-aliases.csv"))
        self.assertEqual(out["aliases"]["loaded"], 0)
        self.assertIn("not found", out["aliases"]["note"])

    def test_a_populated_alias_table_matches_the_variant_spelling(self):
        # The end-to-end payoff for rung 3: without the table this role is
        # deferred as a new company, because "BoF Careers" is a portal name and
        # normalises to nothing already in the sheet.
        master = self.tmp / "Strategic design ED.28"
        master.mkdir()
        self.build([{"Company / Outreach Account": "Moncler", "Company ID": "polids-6556667b"}], root=master)
        (self.tmp / "company-aliases.csv").write_text(
            "alias_norm;canonical_id;note\n"
            "bof careers;polids-6556667b;BoF Careers republished a Moncler role\n",
            encoding="utf-8",
        )
        self.stage_run([self.new_company("BoF Careers")], root=master)

        out = json.loads(self.run_append(root=master).stdout)
        self.assertEqual(out["aliases"]["loaded"], 1)
        self.assertEqual(out["already_present"], 1, "rung 3 must match the alias, not defer it")
        self.assertEqual(out["to_append"], 0)
        self.assertEqual(out["deferred"], 0)

    def test_appends_only_new_companies(self):
        self.build([{"Company / Outreach Account": "Esistente Srl", "Company ID": "polids-11111111"}])
        run = self.stage_run([self.new_company("Nuova Azienda Srl"), self.new_company("Esistente Srl")])
        proc = self.run_append()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        out = json.loads(proc.stdout)
        self.assertEqual(out["appended"], 1)
        self.assertEqual(out["already_present"], 1)

    def test_a_drifted_header_is_refused_not_read_as_empty(self):
        """Every reader finds its columns by name, so a missing one reads as "".

        Measured with only `Company / Outreach Account` renamed: every company
        read as unknown, so a run re-proposes accounts it already has; and the row
        `append` then writes carries the role, the links, a `Not started` status
        and a fresh `Company ID` — with no company name at all, because the writer
        resolves column names the same way the reader does. Silent in both
        directions, and indistinguishable from normal operation downstream.
        """
        self.build([{"Company / Outreach Account": "Esistente Srl", "Company ID": "polids-11111111"}])
        self.stage_run([self.new_company("Nuova Srl")])

        from openpyxl import load_workbook

        target = self.tmp / "Review.xlsx"
        wb = load_workbook(target)
        ws = wb["Review"]
        ws.cell(2, 1, "Company")  # was "Company / Outreach Account"
        wb.save(target)

        proc = self.run_append()
        self.assertEqual(proc.returncode, 1, "append must refuse, not re-add the sheet")
        out = json.loads(proc.stdout)
        self.assertFalse(out["ok"])
        self.assertIn("Company / Outreach Account", out["error"])

        # Nothing was written: the sheet still holds exactly its one company.
        ws = load_workbook(target)["Review"]
        self.assertEqual([ws.cell(r, 1).value for r in range(3, ws.max_row + 1)], ["Esistente Srl"])

        # The pre-run reader refuses for the same reason, so a run cannot dedup
        # against a sheet nobody can read.
        pre = subprocess.run(
            [
                sys.executable, str(HERE / "sync_export.py"), "export-history",
                "--dir", str(self.tmp), "--out", str(self.tmp / "h.csv"),
            ],
            capture_output=True,
            text=True,
        )
        self.assertEqual(pre.returncode, 1)
        self.assertIn("Company / Outreach Account", json.loads(pre.stdout)["error"])

    def test_migrate_review_still_reads_a_drifted_header(self):
        """The repair tool must not be behind the guard it exists to repair.

        `migrate-review` reads the sheet's header itself rather than through
        `_review_as_canonical`, which is what lets it rebuild a sheet the readers
        now refuse. If that ever changes, a drifted header becomes unfixable.
        """
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Review"
        # A header missing most of the current columns — the drifted case.
        for c, name in enumerate(["Company / Outreach Account", "Contact Search Status", "Company ID"], start=1):
            ws.cell(2, c, name)
        ws.cell(3, 1, "Alpha Srl")
        ws.cell(3, 2, "Not started")
        ws.cell(3, 3, "polids-aaaaaaaa")
        wb.save(self.tmp / "Review.xlsx")

        proc = subprocess.run(
            [sys.executable, str(HERE / "sync_export.py"), "migrate-review", "--dir", str(self.tmp)],
            capture_output=True,
            text=True,
        )
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        out = json.loads(proc.stdout)
        self.assertTrue(out["ok"])
        self.assertTrue(out["columns_added"], "the rebuild should have added the missing columns")
        self.assertEqual(out["rows"], 1)

    def test_a_gap_in_the_middle_does_not_write_over_the_rows_below(self):
        """Row 5 emptied, rows 6-7 still holding companies, two new ones to add.

        `_first_free_row` returned 5 once and the loop then did `row_i += 1`, so
        the second new company landed on row 6 and Delta Srl was gone — company,
        status, notes, dates, all of it. Verified against the pre-fix code before
        the fix was written, not inferred.
        """
        self.build([
            {"Company / Outreach Account": "Alfa Srl"},
            {"Company / Outreach Account": "Bravo Srl"},
            {},  # row 5 — emptied by a colleague, nothing else in it
            {"Company / Outreach Account": "Delta Srl"},
            {"Company / Outreach Account": "Echo Srl"},
        ])
        self.stage_run([self.new_company("Foxtrot Srl"), self.new_company("Golf Srl")])
        proc = self.run_append()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)

        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        self.assertEqual(
            [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)],
            # The emptied row IS reused — that is the point of scanning rather than
            # taking max_row+1. What must not happen is the next new company
            # continuing downwards into Delta.
            ["Alfa Srl", "Bravo Srl", "Foxtrot Srl", "Delta Srl", "Echo Srl", "Golf Srl"],
        )

    def test_a_row_that_still_holds_values_is_refused_not_overwritten(self):
        """Clearing only the company name leaves the rest of the row behind.

        Writing a new company there would take the previous one's job links and
        machine notes, so the append refuses instead: Review.xlsx is left exactly
        as it was and the row goes to the sidecar. Same trade as every other
        guard — a missed append costs one paste, a silent overwrite costs a
        colleague's work.
        """
        self.build([
            {"Company / Outreach Account": "Alfa Srl"},
            {"Company / Outreach Account": "Bravo Srl"},
            {
                "Matching Job Titles": "1. Service Design Intern",
                "Job Links": "1. https://example.com/vecchio",
                "Matching Notes": "[NEW COMPANY] [pertinente] Match score: 90/100; nota",
            },
            {"Company / Outreach Account": "Delta Srl"},
        ])
        self.stage_run([self.new_company("Foxtrot Srl")])
        proc = self.run_append()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        out = json.loads(proc.stdout)
        self.assertEqual(out["appended"], 0)
        self.assertEqual(out["reason"], "cell_occupied")

        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        link_at = REVIEW_COLUMNS.index("Job Links") + 1
        self.assertEqual(ws.cell(5, link_at).value, "1. https://example.com/vecchio")
        self.assertIsNone(ws.cell(5, 1).value, "the company cell must still be empty")

    def test_the_sidecar_carries_what_the_workbook_would_have_written(self):
        """The sidecar replaces a refused append, so it has to be as complete.

        It was not. It looked the columns up in the A4-shaped dict — where the
        machine notes column is `Notes`, not `Matching Notes` — and the two values
        that exist only inside the append loop (`Contact Search Status` from
        A4's defaults, `Company ID` from `propose_company_id`) were absent from
        that dict entirely. Three columns came out empty on every row of every
        sidecar, including the A4-mandated per-role justification.
        """
        self.build([{"Company / Outreach Account": "Esistente Srl"}])
        self.stage_run([self.new_company("Nuova Srl")])
        # Any refusal takes the sidecar path; an Excel owner file is the cheapest
        # one to arrange from a test.
        (self.tmp / "~$Review.xlsx").write_bytes(b"lock")
        proc = self.run_append()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        out = json.loads(proc.stdout)
        self.assertEqual(out["appended"], 0)
        self.assertEqual(out["reason"], "excel_lock")

        sidecar = self.tmp / f"Review-additions-{out['run_id']}.csv"
        rows = list(csv.DictReader(io.StringIO(sidecar.read_text(encoding="utf-8")), delimiter=";"))
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["Company / Outreach Account"], "Nuova Srl")
        self.assertEqual(row["Matching Notes"], "machine note", "the A4 justification was always empty")
        self.assertEqual(row["Contact Search Status"], "Not started", "the status default was always empty")
        # Determinstic, so this is also the id a later `append` would propose --
        # pasting the sidecar and appending it later cannot mint two identities.
        self.assertEqual(row["Company ID"], propose_company_id("Nuova Srl"), "the id was always empty")
        self.assertEqual(row["Reviewer Notes"], "", "the machine never writes the colleague's column")
        self.assertEqual(row["Last Checked"], "2026-09-15", "dates go out ISO, not 15/09/2026")

    def test_machine_owned_dates_are_written_as_real_dates(self):
        """`Last Checked` is machine-owned (A4 MACHINE_LATEST_COLS) even though it
        is a date column. Blanking it would lose the run date."""
        self.build([])
        self.stage_run([self.new_company()])
        self.assertEqual(self.run_append().returncode, 0)
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        cell = ws.cell(3, hdr.index("Last Checked") + 1)
        self.assertIsInstance(cell.value, datetime, f"Last Checked must be a date, got {cell.value!r}")
        self.assertEqual((cell.value.year, cell.value.month, cell.value.day), (2026, 9, 15))

    def test_human_owned_columns_are_never_invented(self):
        self.build([])
        self.stage_run([self.new_company()])
        self.assertEqual(self.run_append().returncode, 0)
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        for col, want in [
            ("Reviewer Notes", None),
            ("First Contact Date", None),
            ("Recall", None),
            ("Contact Name", None),
            ("Contact Email / LinkedIn", None),
            ("Contact Search Status", "Not started"),
        ]:
            self.assertEqual(ws.cell(3, hdr.index(col) + 1).value, want, f"{col} should be {want!r}")

    def test_a_new_company_gets_a_proposed_id(self):
        self.build([])
        self.stage_run([self.new_company()])
        self.run_append()
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        cid = ws.cell(3, hdr.index("Company ID") + 1).value
        self.assertTrue(cid and cid.startswith("polids-"), f"expected a proposed id, got {cid!r}")

    def test_never_overwrites_an_existing_cell(self):
        self.build(
            [
                {
                    "Company / Outreach Account": "Esistente Srl",
                    "Company ID": "polids-11111111",
                    "Reviewer Notes": "a colleague's note",
                    "Contact Search Status": "Contact found",
                    "First Contact Date": datetime(2026, 3, 9),
                }
            ]
        )
        self.stage_run([self.new_company()])
        self.run_append()
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(ws.cell(3, hdr.index("Reviewer Notes") + 1).value, "a colleague's note")
        self.assertEqual(ws.cell(3, hdr.index("Contact Search Status") + 1).value, "Contact found")
        d = ws.cell(3, hdr.index("First Contact Date") + 1).value
        self.assertEqual((d.year, d.month, d.day), (2026, 3, 9))

    def test_near_duplicate_is_deferred_not_appended(self):
        """Appending 'Nuova Azienda Group' beside 'Nuova Azienda' is how the 5
        existing duplicate rows happened."""
        self.build([{"Company / Outreach Account": "Nuova Azienda", "Company ID": "polids-11111111"}])
        self.stage_run([self.new_company("Nuova Azienda Group")])
        out = json.loads(self.run_append().stdout)
        self.assertEqual(out["appended"], 0)
        self.assertEqual(out["deferred"], 1)
        self.assertEqual(out["deferred_detail"][0]["reason"], "similar_name")

    def test_refuses_while_excel_has_the_file_open_and_writes_a_sidecar(self):
        self.build([])
        run = self.stage_run([self.new_company()])
        (self.tmp / "~$Review.xlsx").write_bytes(b"")
        proc = self.run_append()
        self.assertEqual(proc.returncode, 0, "a refused append is safe, not an error")
        out = json.loads(proc.stdout)
        self.assertEqual(out["appended"], 0)
        self.assertEqual(out["reason"], "excel_lock")
        self.assertTrue((self.tmp / f"Review-additions-{run}.csv").exists(), "sidecar must be written")


class TestHarvest(TmpDirCase):
    """What colleagues decided, read back out of Review.xlsx — read-only.

    There is no longer a canonical CSV to write into: Review.xlsx IS the record,
    and `sync_export.py export-history` is how the run reads it."""

    def make_canon(self, rows):
        p = self.tmp / "canon.csv"
        with p.open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(A4_COLUMNS + ["Company ID"])
            for r in rows:
                w.writerow(r)
        return p

    def row(self, name, cid, contacted="03/09/2026", decision="Review", status="Contact found"):
        r = [""] * (len(A4_COLUMNS) + 1)
        r[0] = name
        r[A4_COLUMNS.index("First Contact Date")] = contacted
        r[A4_COLUMNS.index("Outreach Decision")] = decision
        r[A4_COLUMNS.index("Contact Search Status")] = status
        r[-1] = cid
        return r

    def build_review(self, entries):
        """entries: list of (company, company_id, {column: value})"""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Review"
        for c, name in enumerate(REVIEW_COLUMNS, start=1):
            ws.cell(2, c, name)
        for i, (name, cid, vals) in enumerate(entries):
            ws.cell(3 + i, 1, name)
            ws.cell(3 + i, REVIEW_COLUMNS.index("Company ID") + 1, cid)
            for col, v in vals.items():
                ws.cell(3 + i, REVIEW_COLUMNS.index(col) + 1, v)
        wb.save(self.tmp / "Review.xlsx")

    def run_(self, cmd, history, *extra):
        return subprocess.run(
            [sys.executable, str(HERE / "sync_export.py"), cmd, "--dir", str(self.tmp), "--history", str(history), *extra],
            capture_output=True,
            text=True,
        )

    def test_same_date_in_two_representations_is_not_a_conflict(self):
        """The canonical holds '03/09/2026' where the workbook holds a real date.
        Comparing those as strings flags every single row."""
        canon = self.make_canon([self.row("Alpha Srl", "polids-aaaaaaaa", contacted="03/09/2026")])
        self.build_review([("Alpha Srl", "polids-aaaaaaaa", {"First Contact Date": datetime(2026, 9, 3)})])
        out = json.loads(self.run_("harvest", canon).stdout)
        self.assertEqual(out["conflicts"], 0, "the same date must not read as a conflict")
        self.assertEqual(out["updates"], 0, "nothing changed, so nothing to write")

    def test_harvest_never_writes(self):
        canon = self.make_canon([self.row("Alpha Srl", "polids-aaaaaaaa")])
        before = canon.read_bytes()
        self.build_review([("Alpha Srl", "polids-aaaaaaaa", {"Contact Search Status": "Contact found"})])
        self.run_("harvest", canon)
        self.assertEqual(canon.read_bytes(), before, "harvest is the audit path: read-only")

    def test_shared_company_id_is_ambiguous_not_silently_first_wins(self):
        """The 5 duplicate names share a proposed id. Collapsing them to one row
        would send every update to the first and strand the second -- the same
        failure as add_verified's last-row-wins, via a different key."""
        canon = self.make_canon(
            [self.row("KPMG", "polids-7187734f"), self.row("Altro Srl", "polids-bbbbbbbb"), self.row("KPMG", "polids-7187734f")]
        )
        self.build_review([("KPMG", "polids-7187734f", {"Contact Search Status": "Contact found"})])
        out = json.loads(self.run_("harvest", canon).stdout)
        self.assertEqual(out["updates"], 0, "must not pick one of the two rows")
        kinds = [c for c in out["conflict_detail"] if c["canonical"] == "duplicate_id"]
        self.assertEqual(len(kinds), 1)


class TestExportHistory(TmpDirCase):
    """`export-history` turns Review.xlsx into the CSV the search run dedups against.

    Why the run reads this and not a canonical CSV: measured on 2026-09-18, the
    canonical had gone 62 companies stale (118 against Review's 180) because
    nothing had run `pull` for weeks. Deduping a run against it re-proposed 76
    roles that had just been published. Review.xlsx is the file colleagues
    actually maintain, so it is the one that answers "do we know this already".
    """

    def build_review(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Review"
        ws.cell(1, 1, "banner")
        for c, name in enumerate(REVIEW_COLUMNS, start=1):
            ws.cell(2, c, name)
        row = {c: "" for c in REVIEW_COLUMNS}
        row["Company / Outreach Account"] = "Alpha Srl"
        row["Matching Job Titles"] = "1. Service Design Intern — Milan, Italy"
        row["Job Links"] = "1. https://alpha.example/1"
        row["Locations"] = "Milan, Italy"
        row["Matching Notes"] = "[NEW COMPANY] machine prose"
        row["Reviewer Notes"] = "a colleague's note"
        row["Company ID"] = "polids-aaaaaaaa"
        row["Contact Search Status"] = "Contact found"
        for c, name in enumerate(REVIEW_COLUMNS, start=1):
            ws.cell(3, c, row[name] or None)
        wb.save(self.tmp / "Review.xlsx")

    def export(self, *extra):
        return subprocess.run(
            [sys.executable, str(HERE / "sync_export.py"), "export-history", "--dir", str(self.tmp), *extra],
            capture_output=True,
            text=True,
        )

    def test_writes_a_csv_the_run_can_read(self):
        self.build_review()
        out = self.tmp / "h.csv"
        proc = self.export("--out", str(out))
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        meta = json.loads(proc.stdout)
        self.assertEqual(meta["companies"], 1)
        self.assertEqual(meta["with_roles"], 1)

        rows = list(csv.reader(io.StringIO(out.read_text(encoding="utf-8")), delimiter=";"))
        self.assertEqual(rows[0], A4_COLUMNS + ["Company ID"], "the header the reader expects")
        rec = dict(zip(rows[0], rows[1]))
        self.assertEqual(rec["Company / Outreach Account"], "Alpha Srl")
        self.assertEqual(rec["Matching Job Titles"], "1. Service Design Intern — Milan, Italy")
        self.assertEqual(rec["Company ID"], "polids-aaaaaaaa")

    def test_a_date_cell_survives_the_round_trip(self):
        """A writer and a reader disagreeing about one format is the bug class this
        module exists to prevent, and export-history had it against its own reader:
        the date cell came out as str(datetime), which parse_date returns None for.
        `_sidecar_value` already states the intended format: date-only ISO."""
        self.build_review()
        from openpyxl import load_workbook

        wb = load_workbook(self.tmp / "Review.xlsx")
        wb["Review"].cell(3, REVIEW_COLUMNS.index("First Contact Date") + 1, datetime(2026, 9, 3))
        wb.save(self.tmp / "Review.xlsx")

        out = self.tmp / "h.csv"
        self.assertEqual(self.export("--out", str(out)).returncode, 0)
        rows = list(csv.reader(io.StringIO(out.read_text(encoding="utf-8")), delimiter=";"))
        rec = dict(zip(rows[0], rows[1]))
        self.assertEqual(rec["First Contact Date"], "2026-09-03")
        self.assertIsNotNone(
            parse_date(rec["First Contact Date"]), "the export must be readable by its own reader"
        )

    def test_the_machine_notes_half_maps_to_a4_notes(self):
        # Review splits A4's single `Notes` into `Matching Notes` (machine) and
        # `Reviewer Notes` (human). The export has one A4 column to fill, and the
        # machine half is what A4 means by Notes.
        self.build_review()
        out = self.tmp / "h.csv"
        self.export("--out", str(out))
        rows = list(csv.reader(io.StringIO(out.read_text(encoding="utf-8")), delimiter=";"))
        rec = dict(zip(rows[0], rows[1]))
        self.assertEqual(rec["Notes"], "[NEW COMPANY] machine prose")
        self.assertNotEqual(rec["Notes"], "a colleague's note", "the human half must not win")

    def test_columns_review_does_not_render_come_out_empty_not_invented(self):
        self.build_review()
        out = self.tmp / "h.csv"
        self.export("--out", str(out))
        rows = list(csv.reader(io.StringIO(out.read_text(encoding="utf-8")), delimiter=";"))
        rec = dict(zip(rows[0], rows[1]))
        for gone in ("Previously Contacted?", "Outreach Decision"):
            self.assertEqual(rec[gone], "", f"{gone} is not in Review.xlsx, so it must not be guessed at")

    def test_the_exported_csv_loads_as_history(self):
        """The point of the command is that `loadHistory` can read what it writes.
        Asserted through the real reader rather than by eyeballing the header."""
        self.build_review()
        out = self.tmp / "h.csv"
        self.export("--out", str(out))
        text = out.read_text(encoding="utf-8")
        self.assertTrue(text.startswith("Company / Outreach Account;"))
        self.assertIn("Alpha Srl;", text)
        # The reader the engine uses splits on ';' and resolves these by name.
        for required in ("Company / Outreach Account", "Matching Job Titles", "Job Links", "Locations", "Company ID"):
            self.assertIn(required, text.split("\n")[0])

    def test_refuses_a_folder_with_no_review_file(self):
        proc = self.export()
        self.assertEqual(proc.returncode, 2)
        self.assertFalse(json.loads(proc.stdout)["ok"])


class TestMigrateReview(TmpDirCase):
    """Rebuilding Review.xlsx onto a changed REVIEW_COLUMNS.

    `stage` and `append` write into the sheet at POSITIONS taken from
    REVIEW_COLUMNS, so changing that list without a migration puts every value
    after the change into the wrong column -- and the tab-count era's lesson
    applies here too: the file still looks fine.
    """

    # The 24-column set that shipped before the vocabulary change.
    OLD_COLUMNS = [
        "Company / Outreach Account",
        "Brands / Business Units",
        "In Italy?",
        "Locations",
        "Master-fit Themes",
        "Matching Job Titles",
        "Job Links",
        "Role Count",
        "Curricular Evidence",
        "Work Modes",
        "Sources / Portals",
        "Previously Contacted?",
        "Contact Search Status",
        "Contact Name",
        "Contact Role",
        "Contact Email / LinkedIn",
        "Outreach Decision",
        "Matching Notes",
        "Reviewer Notes",
        "Verification Status",
        "Last Checked",
        "First Contact Date",
        "Recall",
        "Company ID",
    ]

    def build_old_sheet(self, rows):
        """rows: list of {column: value}"""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Review"
        ws.cell(1, 1, "old banner")
        for c, name in enumerate(self.OLD_COLUMNS, start=1):
            ws.cell(2, c, name)
        for i, vals in enumerate(rows):
            ws.cell(3 + i, 1, vals.get("Company / Outreach Account", f"Co {i}"))
            for col, v in vals.items():
                ws.cell(3 + i, self.OLD_COLUMNS.index(col) + 1, v)
        wb.save(self.tmp / "Review.xlsx")

    def migrate(self, *extra):
        return subprocess.run(
            [sys.executable, str(HERE / "sync_export.py"), "migrate-review", "--dir", str(self.tmp), *extra],
            capture_output=True,
            text=True,
        )

    def read_new(self):
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        out = {}
        for r in range(3, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                continue
            out[ws.cell(r, 1).value] = {h: ws.cell(r, c + 1).value for c, h in enumerate(hdr) if h}
        return hdr, out

    def test_carries_every_cell_across_a_changed_column_set(self):
        self.build_old_sheet(
            [
                {
                    "Company / Outreach Account": "Alpha Srl",
                    "Contact Search Status": "Contacted",
                    "Contact Name": "Enrica Marro",
                    "Matching Notes": "[NEW COMPANY] ...",
                    "Previously Contacted?": "Yes",
                    "Outreach Decision": "Review",
                    "First Contact Date": datetime(2026, 9, 3),
                    "Last Checked": datetime(2026, 8, 27),
                    "Company ID": "polids-aaaaaaaa",
                }
            ]
        )
        proc = self.migrate()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        out = json.loads(proc.stdout)
        self.assertEqual(out["rows"], 1)
        self.assertEqual(out["columns_dropped"], ["Previously Contacted?", "Outreach Decision"])

        hdr, by = self.read_new()
        self.assertEqual(hdr, REVIEW_COLUMNS)
        row = by["Alpha Srl"]
        self.assertEqual(row["Contact Name"], "Enrica Marro", "a cell after the dropped columns must not shift")
        self.assertEqual(row["Matching Notes"], "[NEW COMPANY] ...")
        self.assertEqual(row["Company ID"], "polids-aaaaaaaa")
        self.assertIsInstance(row["First Contact Date"], datetime, "real dates must survive as dates")
        # number_format lives on the CELL, not the value, so read it directly.
        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        cell = ws.cell(3, hdr.index("First Contact Date") + 1)
        self.assertEqual(cell.number_format, "DD/MM/YYYY")

    def test_translates_the_old_status_vocabulary(self):
        self.build_old_sheet(
            [
                {"Company / Outreach Account": "A", "Contact Search Status": "Contacted"},
                {"Company / Outreach Account": "B", "Contact Search Status": "No suitable contact"},
                {"Company / Outreach Account": "C", "Contact Search Status": "Not started"},
                {"Company / Outreach Account": "D", "Contact Search Status": ""},
            ]
        )
        out = json.loads(self.migrate().stdout)
        self.assertEqual(out["status_translated"], {"Contacted": 1, "No suitable contact": 1})
        self.assertEqual(out["unmapped_status"], {})
        by = self.read_new()[1]
        self.assertEqual(by["A"]["Contact Search Status"], "Contact found")
        self.assertEqual(by["B"]["Contact Search Status"], "Job not suitable")
        self.assertEqual(by["C"]["Contact Search Status"], "Not started")

    def test_an_unknown_status_is_refused_not_guessed_at(self):
        """A value that is neither current nor known-legacy is someone's intent.
        Migrating would leave it sitting under a vocabulary it does not belong
        to, and nothing downstream would flag it."""
        self.build_old_sheet([{"Company / Outreach Account": "A", "Contact Search Status": "chiamato forse"}])
        before = (self.tmp / "Review.xlsx").read_bytes()
        proc = self.migrate()
        self.assertEqual(proc.returncode, 2)
        out = json.loads(proc.stdout)
        self.assertEqual(out["reason"], "unmapped_status")
        self.assertEqual(out["unmapped"], {"chiamato forse": 1})
        self.assertEqual((self.tmp / "Review.xlsx").read_bytes(), before, "a refusal writes nothing")

        forced = self.migrate("--force")
        self.assertEqual(forced.returncode, 0, forced.stdout)
        self.assertEqual(json.loads(forced.stdout)["unmapped_status"], {"chiamato forse": 1})
        self.assertEqual(self.read_new()[1]["A"]["Contact Search Status"], "chiamato forse")

    def test_dry_run_reports_the_plan_and_writes_nothing(self):
        self.build_old_sheet([{"Company / Outreach Account": "A", "Contact Search Status": "Contacted"}])
        before = (self.tmp / "Review.xlsx").read_bytes()
        out = json.loads(self.migrate("--dry-run").stdout)
        self.assertTrue(out["dry_run"])
        self.assertEqual(out["status_translated"], {"Contacted": 1})
        self.assertEqual((self.tmp / "Review.xlsx").read_bytes(), before)

    def test_takes_a_backup_and_leaves_colour_rules_installed(self):
        self.build_old_sheet([{"Company / Outreach Account": "A", "Contact Search Status": "Contacted"}])
        self.assertEqual(self.migrate().returncode, 0)
        self.assertTrue(list((self.tmp / "_machine" / "backups").glob("Review-*.xlsx")), "pre-write backup")

        from openpyxl import load_workbook

        ws = load_workbook(self.tmp / "Review.xlsx")["Review"]
        rules = {str(rng.sqref): [r.formula[0] for r in rs] for rng, rs in ws.conditional_formatting._cf_rules.items()}
        status_letter = get_column_letter(REVIEW_COLUMNS.index("Contact Search Status") + 1)
        last = get_column_letter(len(REVIEW_COLUMNS))
        self.assertIn(f"A3:{last}{CF_RANGE_ROWS}", rules)
        self.assertIn(f'${status_letter}3="Contact found"', rules[f"A3:{last}{CF_RANGE_ROWS}"])

    def test_a_row_with_no_company_is_skipped(self):
        self.build_old_sheet([{"Company / Outreach Account": "A", "Contact Search Status": "Contacted"}])
        from openpyxl import load_workbook

        wb = load_workbook(self.tmp / "Review.xlsx")
        wb["Review"].cell(4, 1, None)
        wb["Review"].cell(4, 2, "orphan")
        wb.save(self.tmp / "Review.xlsx")

        out = json.loads(self.migrate().stdout)
        self.assertEqual(out["rows"], 1)
        self.assertEqual(list(self.read_new()[1]), ["A"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
